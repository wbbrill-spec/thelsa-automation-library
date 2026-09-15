"""
sit.py — SIT's "Plan de Viajes" (the national trip plan) as a third source.

SIT plans every Mexican domestic truck movement in one macro-enabled workbook
(form THFOPL04, refreshed ~3x a day). Two sheets matter:

  "Plan de Viajes"        one row per movement: executive, origin, destination,
                          customer, service type, load/unload day + hour, m³,
                          branch, unit (truck), driver, freight, km, notes.
  "Programación Unidades" the fleet: one row per unit with its driver, home
                          plaza, CAP (usable m³, 85 on every unit today) and a
                          day-by-day load/unload schedule for the week.

Why the cross-border dashboard cares
------------------------------------
Thelsa's own cross-border moves already ride SIT trucks: rows whose EJECUTIVO
is `TIM/<name>` or `TMS/<name>` are TIM and TMS files, ~50 of 459 on the
2026-09-11 plan. That is where the *onward leg* lives — the truck, driver and
dates between the border/hub and the customer — which neither ClickUp nor
Moveware holds. It also gives the consolidation engine real trucks with real
spare capacity (CAP minus booked m³) instead of a hypothetical 88 m³ trailer.

Matching (measured on the 2026-09-11 plan)
------------------------------------------
* TMS rows: the column headed TELEFONO actually carries the **Moveware job
  number** (110719, 110801, 111127 … sometimes with a letter suffix like
  109446A) — an exact key. 34 of 48 Thelsa rows carried one.
* TIM rows: no reference at all (`S/N`, `.`, blank), so they match on the
  customer name, the same reference-then-fuzzy-name approach `remisiones.py`
  uses.
* One job spans several rows: a line haul (FOR / TRAN F COMP) plus its
  services (DESEMPAQUE, HUACAL, SUELTO, MUEBLES), and a big job can split
  across several units on the same day. Rows are therefore grouped into one
  `JobPlan` per job, with the line-haul legs kept separate from services.

Read-only. This module never writes to the workbook.

Where the workbook lives
------------------------
Sergio Garza Padilla's OneDrive → **Documents / Plan de Viajes**. Operations
does not overwrite one file: each refresh is a NEW workbook named for the day
and the pass — "PLAN DE VIAJES 11 DE SEPTIEMBRE 2026 (1).xlsm", "(2)", "(3)" —
landing about 15:00, 19:30 and 22:20 UTC (38 files over the 13 days to
2026-09-11). The reader therefore lists the folder and takes the most recently
modified .xlsm, which also survives a missed or extra pass.

Env vars
--------
  SIT_PLAN_DRIVE_ID / SIT_PLAN_FOLDER_ID   the OneDrive drive + "Plan de Viajes"
      folder ids (defaults below, read from the share link on 2026-09-11).
  SIT_PLAN_ITEM_ID                       pin one specific workbook (testing).
  SIT_PLAN_XLSX_PATH                     optional local path (dev / manual drop).
  SIT_UNIT_CAP_M3                        fallback capacity when the fleet sheet
                                         has none (default 85, SIT's own figure).
"""
from __future__ import annotations

import datetime as dt
import difflib
import io
import os
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Optional

from .models import Hub, Source, hub_for_destination, norm_text, parse_date, to_number

PLAN_SHEET = "Plan de Viajes"
FLEET_SHEET = "Programación Unidades"

# Column order of the "Plan de Viajes" sheet (header on row 3, data from row 4).
COLUMNS = [
    "ejecutivo", "origen", "desvio", "destino", "cliente", "referencia", "tipo",
    "dia_carga", "hora_carga", "dia_descarga", "hora_descarga", "auto", "m3",
    "sucursal", "suc_apoyo", "suc_apoyo_destino", "jc_descarga", "unidad",
    "operador", "dia_solicitado", "flete", "empresa", "mot_cancelacion", "km",
    "total_servicio", "notas", "consecutivo",
]
HEADER_ROW = 3
FIRST_DATA_ROW = 4

# Service types that are an actual truck movement rather than a service at
# either end. Normalized (accent- and space-stripped) before lookup, which also
# folds the spelling variants operations types by hand: TRAN F COMP /
# TRAN FCOMP / TRANS F COMP / TRAN FCOM, DESEMPAQUE / DESEPAQUE.
LINE_HAUL_TYPES = {"for", "cargafor", "tranfcomp", "transfcomp", "tranfcom",
                   "entfcomp", "intc", "intd", "loc"}
SERVICE_TYPES = {"desempaque", "desepaque", "huacal", "suelto", "muebles"}

DEFAULT_UNIT_CAP_M3 = 85.0

# Sergio Garza Padilla's OneDrive → Documents / Plan de Viajes, resolved from the
# share link Bill supplied on 2026-09-11. Neither id is a secret.
DEFAULT_DRIVE_ID = "b!tOzEkUhjPkW10ykKgNmGkU6ERtSaUo1On4YJxSSIYF8_w2a2XMJAQImqRT51x1Vd"
DEFAULT_FOLDER_ID = "01G6OA2VPRONL6W3YUVJGKX2LBTM6UMFDK"

# A Moveware job number as it appears in the plan: six digits, optionally with a
# trailing letter for a split file (109446A).
_JOB_RE = re.compile(r"^(1\d{5})([A-Za-z])?$")


def _clean(v) -> str:
    """Trim, collapse whitespace and repair the mojibake the workbook carries
    (it is written by a Windows tool and round-trips through cp1252, so names
    arrive as 'MarÃ­a' / 'CUAUHTÃ‰MOC')."""
    if v is None:
        return ""
    s = str(v)
    if "Ã" in s or "Â" in s:
        try:
            s = s.encode("cp1252", errors="strict").decode("utf-8", errors="strict")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    s = unicodedata.normalize("NFC", s)
    return re.sub(r"\s+", " ", s).strip()


def _type_key(v) -> str:
    return re.sub(r"[^a-z]", "", norm_text(v))


def job_number(v) -> str:
    """The Moveware job number carried in the TELEFONO column, or ''."""
    s = _clean(v).replace(" ", "")
    if not s or s.upper() in {"S/N", "SN", "N/A", "NA", "."}:
        return ""
    m = _JOB_RE.match(s)
    return (m.group(1) + (m.group(2) or "").upper()) if m else ""


@dataclass
class Trip:
    """One row of the plan."""
    row: int
    ejecutivo: str = ""
    source: Optional[Source] = None    # TIM / TMS when the row is Thelsa's
    coordinator: str = ""              # the name after TIM/ or TMS/
    job_no: str = ""                   # Moveware job number (TMS rows)
    customer: str = ""
    origin: str = ""
    destination: str = ""
    tipo: str = ""
    load_date: Optional[dt.date] = None
    unload_date: Optional[dt.date] = None
    m3: float = 0.0
    branch: str = ""
    dest_branch: str = ""
    unit: str = ""
    driver: str = ""
    notes: str = ""

    @property
    def is_line_haul(self) -> bool:
        return _type_key(self.tipo) in LINE_HAUL_TYPES

    @property
    def is_thelsa(self) -> bool:
        return self.source is not None

    @property
    def scheduled(self) -> bool:
        """A truck and a day are both assigned."""
        return bool(self.unit and self.load_date)

    def to_dict(self) -> dict:
        return {
            "row": self.row, "source": self.source.value if self.source else None,
            "coordinator": self.coordinator, "job_no": self.job_no,
            "customer": self.customer, "origin": self.origin,
            "destination": self.destination, "tipo": self.tipo,
            "load_date": self.load_date.isoformat() if self.load_date else None,
            "unload_date": self.unload_date.isoformat() if self.unload_date else None,
            "m3": self.m3, "branch": self.branch, "dest_branch": self.dest_branch,
            "unit": self.unit, "driver": self.driver, "notes": self.notes,
            "line_haul": self.is_line_haul,
        }


@dataclass
class Unit:
    """One truck from the fleet sheet."""
    code: str
    driver: str = ""
    plaza: str = ""
    tipo: str = ""
    capacity_m3: float = DEFAULT_UNIT_CAP_M3

    def to_dict(self) -> dict:
        return {"code": self.code, "driver": self.driver, "plaza": self.plaza,
                "tipo": self.tipo, "capacity_m3": self.capacity_m3}


@dataclass
class JobPlan:
    """Every plan row belonging to one Thelsa job."""
    key: str                            # job number, else a normalized name
    source: Optional[Source] = None
    job_no: str = ""
    customer: str = ""
    coordinator: str = ""
    trips: list[Trip] = field(default_factory=list)

    @property
    def line_hauls(self) -> list[Trip]:
        return [t for t in self.trips if t.is_line_haul]

    @property
    def services(self) -> list[Trip]:
        return [t for t in self.trips if not t.is_line_haul]

    @property
    def units(self) -> list[str]:
        seen: list[str] = []
        for t in self.line_hauls or self.trips:
            if t.unit and t.unit not in seen:
                seen.append(t.unit)
        return seen

    @property
    def next_load(self) -> Optional[Trip]:
        """The earliest scheduled line haul — what the board should show."""
        dated = sorted((t for t in self.line_hauls if t.load_date),
                       key=lambda t: t.load_date)
        return dated[0] if dated else None

    @property
    def m3(self) -> float:
        return round(sum(t.m3 for t in self.line_hauls), 2)

    def to_dict(self) -> dict:
        nl = self.next_load
        return {
            "key": self.key, "source": self.source.value if self.source else None,
            "job_no": self.job_no, "customer": self.customer,
            "coordinator": self.coordinator, "units": self.units, "m3": self.m3,
            "trip_count": len(self.trips), "line_haul_count": len(self.line_hauls),
            "next_load": nl.to_dict() if nl else None,
            "awaiting_truck": bool(self.line_hauls) and not any(t.scheduled for t in self.line_hauls),
            "trips": [t.to_dict() for t in self.trips],
        }


# ── parsing ──────────────────────────────────────────────────────────────────
def parse_trip_row(values: list, row_no: int) -> Optional[Trip]:
    """One sheet row → Trip, or None when the row is blank//a spacer."""
    d = dict(zip(COLUMNS, list(values) + [None] * len(COLUMNS)))
    ejec = _clean(d["ejecutivo"])
    origin = _clean(d["origen"])
    if not ejec and not origin:
        return None
    src = coord = ""
    head = ejec.upper()
    if head.startswith("TIM/") or head.startswith("TMS/"):
        src, coord = head[:3], ejec[4:].strip()
    m3 = to_number(d["m3"]) or 0.0
    return Trip(
        row=row_no, ejecutivo=ejec,
        source=Source(src) if src else None, coordinator=coord,
        job_no=job_number(d["referencia"]),
        customer=_clean(d["cliente"]), origin=origin,
        destination=_clean(d["destino"]), tipo=_clean(d["tipo"]),
        load_date=parse_date(d["dia_carga"]), unload_date=parse_date(d["dia_descarga"]),
        m3=round(float(m3), 3), branch=_clean(d["sucursal"]),
        dest_branch=_clean(d["suc_apoyo_destino"]), unit=_clean(d["unidad"]).upper(),
        driver=_clean(d["operador"]), notes=_clean(d["notas"]),
    )


def parse_plan_sheet(rows: list[list]) -> list[Trip]:
    """`rows` are the sheet's rows from row 1; the header sits on row 3."""
    out = []
    for i, values in enumerate(rows, start=1):
        if i < FIRST_DATA_ROW:
            continue
        t = parse_trip_row(values, i)
        if t:
            out.append(t)
    return out


def _cap(v) -> float:
    n = to_number(v)
    return float(n) if n and n > 0 else 0.0


def parse_fleet_sheet(rows: list[list]) -> dict[str, Unit]:
    """"Programación Unidades" → {unit code: Unit}. The header row carries
    OPERADOR / PLAZA / TIPO / CAP / ZONA / ECO, but the columns shift by one on
    rows where TIPO is blank, so each row is read by locating its unit code."""
    fallback = _cap(os.environ.get("SIT_UNIT_CAP_M3")) or DEFAULT_UNIT_CAP_M3
    units: dict[str, Unit] = {}
    for values in rows:
        cells = [_clean(v) for v in list(values)[:8]]
        if not any(cells):
            continue
        eco = next((c for c in cells if re.fullmatch(r"[A-Z]{1,2}\d{1,3}", c.upper())), "")
        if not eco or eco.upper() in units:
            continue
        nums = [_cap(c) for c in cells]
        capacity = next((n for n in nums if 10 <= n <= 200), fallback)
        # A row reads: OPERADOR | PLAZA | [TIPO] | CAP | ECO — TIPO is sometimes
        # blank, so identify the type first and take the plaza from what's left
        # (MTY, GDL, VHSA, "MC. ALLEN" …), never the eco code or a number.
        tail = cells[1:]
        tipo = next((c for c in tail if _type_key(c) in LINE_HAUL_TYPES), "")
        plaza = next((c for c in tail
                      if c != tipo and c.upper() != eco.upper()
                      and not _cap(c) and re.fullmatch(r"[A-Za-z. ]{2,12}", c)), "")
        units[eco.upper()] = Unit(
            code=eco.upper(), driver=cells[0] if cells and not cells[0].isdigit() else "",
            plaza=plaza.upper(), tipo=tipo.upper(), capacity_m3=capacity,
        )
    return units


def load_workbook_trips(data: bytes) -> tuple[list[Trip], dict[str, Unit], dict]:
    """Parse the workbook bytes into trips + fleet, with diagnostics."""
    import openpyxl
    info: dict = {}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        return [], {}, {"error": f"workbook unreadable: {type(exc).__name__}: {exc}"}
    names = wb.sheetnames
    info["sheets"] = names
    plan = next((n for n in names if norm_text(n) == norm_text(PLAN_SHEET)), None)
    if not plan:
        return [], {}, {**info, "error": f"no '{PLAN_SHEET}' sheet"}
    trips = parse_plan_sheet([list(r) for r in wb[plan].iter_rows(
        min_col=1, max_col=len(COLUMNS), values_only=True)])
    fleet_name = next((n for n in names if norm_text(n) == norm_text(FLEET_SHEET)), None)
    fleet = parse_fleet_sheet([list(r) for r in wb[fleet_name].iter_rows(
        min_col=1, max_col=8, values_only=True)]) if fleet_name else {}
    info.update(trips=len(trips), units=len(fleet),
                thelsa_trips=sum(1 for t in trips if t.is_thelsa))
    return trips, fleet, info


# ── grouping Thelsa's rows into jobs ─────────────────────────────────────────
def job_plans(trips: list[Trip]) -> list[JobPlan]:
    """Group the TIM/TMS rows into one JobPlan per job (job number when there
    is one, else the customer name)."""
    by_key: dict[str, JobPlan] = {}
    for t in trips:
        if not t.is_thelsa:
            continue
        key = t.job_no or f"name:{norm_text(t.customer)}"
        if not key or key == "name:":
            continue
        jp = by_key.get(key)
        if jp is None:
            jp = by_key[key] = JobPlan(key=key, source=t.source, job_no=t.job_no,
                                       customer=t.customer, coordinator=t.coordinator)
        jp.trips.append(t)
        if not jp.customer:
            jp.customer = t.customer
    return list(by_key.values())


# ── matching plans to dashboard shipments ────────────────────────────────────
NAME_THRESHOLD = 0.82


def _ref_keys(s) -> set[str]:
    """Every token of a shipment that could appear as a job number."""
    keys = set()
    for raw in (getattr(s, "reference_number", ""), getattr(s, "source_ref", "")):
        for tok in re.split(r"[^0-9A-Za-z]+", str(raw or "")):
            m = _JOB_RE.match(tok)
            if m:
                keys.add(m.group(1) + (m.group(2) or "").upper())
                keys.add(m.group(1))
    return keys


def match_plans_to_shipments(plans: list[JobPlan], shipments: list) -> dict:
    """Attach each JobPlan to a shipment. Job number first (exact, TMS), then
    customer name above NAME_THRESHOLD within the same source.

    Returns {"matches": {shipment_id: JobPlan}, "diag": {...}}. Pure.
    """
    by_ref: dict[str, list] = {}
    for s in shipments or []:
        for k in _ref_keys(s):
            by_ref.setdefault(k, []).append(s)

    matches: dict[str, JobPlan] = {}
    how: dict[str, str] = {}
    unmatched: list[str] = []
    for jp in plans:
        hit = None
        if jp.job_no:
            cands = by_ref.get(jp.job_no) or by_ref.get(jp.job_no.rstrip("ABCDEFGHIJ"))
            same = [s for s in (cands or []) if s.source is jp.source] or (cands or [])
            if len(same) == 1:
                hit, why = same[0], "reference"
            elif len(same) > 1:
                unmatched.append(f"{jp.key} (reference {jp.job_no} on {len(same)} shipments)")
                continue
        if hit is None and jp.customer:
            pool = [s for s in (shipments or []) if s.source is jp.source]
            best, score = None, 0.0
            target = norm_text(jp.customer)
            for s in pool:
                r = difflib.SequenceMatcher(None, target, norm_text(s.customer_name)).ratio()
                if r > score:
                    best, score = s, r
            if best is not None and score >= NAME_THRESHOLD:
                hit, why = best, f"name {score:.2f}"
        if hit is None:
            unmatched.append(jp.key)
            continue
        if hit.id in matches:                       # keep the richer plan
            if len(jp.trips) <= len(matches[hit.id].trips):
                continue
        matches[hit.id] = jp
        how[hit.id] = why
    return {"matches": matches,
            "diag": {"plans": len(plans), "matched": len(matches),
                     "by_reference": sum(1 for v in how.values() if v == "reference"),
                     "by_name": sum(1 for v in how.values() if v.startswith("name")),
                     "unmatched": unmatched[:20], "unmatched_count": len(unmatched)}}


def enrich_shipment(s, jp: JobPlan) -> None:
    """Hang the plan on the shipment for the board and the drawer."""
    nl = jp.next_load
    s.extra = dict(s.extra or {})
    s.extra["sit"] = jp.to_dict()
    if nl:
        s.extra["sit_unit"] = nl.unit
        s.extra["sit_driver"] = nl.driver
        s.extra["sit_load_date"] = nl.load_date.isoformat() if nl.load_date else None
        s.extra["sit_unload_date"] = nl.unload_date.isoformat() if nl.unload_date else None
        s.extra["sit_route"] = f"{nl.origin} → {nl.destination}" if nl.origin else nl.destination
    if jp.to_dict()["awaiting_truck"]:
        if "awaiting_truck" not in s.status_flags:
            s.status_flags = list(s.status_flags) + ["awaiting_truck"]


# ── truck load / spare capacity (input for the consolidation engine) ─────────
def truck_loads(trips: list[Trip], fleet: dict[str, Unit],
                today: Optional[dt.date] = None, horizon_days: int = 14) -> list[dict]:
    """Booked m³ per unit per load-day over the horizon, with the spare space
    left against that unit's capacity. This is what lets the engine propose a
    real truck instead of a hypothetical trailer."""
    today = today or dt.date.today()
    end = today + dt.timedelta(days=horizon_days)
    buckets: dict[tuple[str, dt.date], dict] = {}
    for t in trips:
        if not (t.is_line_haul and t.unit and t.load_date):
            continue
        if not (today <= t.load_date <= end):
            continue
        b = buckets.setdefault((t.unit, t.load_date), {
            "unit": t.unit, "date": t.load_date, "booked_m3": 0.0, "trips": [],
            "capacity_m3": (fleet.get(t.unit).capacity_m3 if fleet.get(t.unit)
                            else DEFAULT_UNIT_CAP_M3),
            "driver": (fleet.get(t.unit).driver if fleet.get(t.unit) else t.driver),
        })
        b["booked_m3"] = round(b["booked_m3"] + t.m3, 2)
        b["trips"].append(t.to_dict())
    out = []
    for b in buckets.values():
        spare = round(b["capacity_m3"] - b["booked_m3"], 2)
        out.append({**b, "date": b["date"].isoformat(), "spare_m3": max(spare, 0.0),
                    "fill_pct": round(b["booked_m3"] / b["capacity_m3"] * 100)
                    if b["capacity_m3"] else 0,
                    "lanes": sorted({f"{t['branch'] or t['origin']} → {t['dest_branch'] or t['destination']}"
                                     for t in b["trips"]})})
    return sorted(out, key=lambda x: (x["date"], x["unit"]))


# ── fetching the live workbook ───────────────────────────────────────────────
def fetch_workbook_bytes() -> tuple[Optional[bytes], dict]:
    """Local path first (dev), then SharePoint/OneDrive via the app-only Graph
    token — the same credential and the same Files.Read.All requirement as
    `remisiones.py`."""
    path = os.environ.get("SIT_PLAN_XLSX_PATH", "").strip()
    if path and os.path.exists(path):
        with open(path, "rb") as fh:
            return fh.read(), {"source": "file", "path": path}
    drive = os.environ.get("SIT_PLAN_DRIVE_ID", DEFAULT_DRIVE_ID).strip()
    folder = os.environ.get("SIT_PLAN_FOLDER_ID", DEFAULT_FOLDER_ID).strip()
    item = os.environ.get("SIT_PLAN_ITEM_ID", "").strip()   # pin one file (testing)
    if not drive or not (folder or item):
        return None, {"source": "graph",
                      "error": "SIT_PLAN_DRIVE_ID plus SIT_PLAN_FOLDER_ID (or "
                               "SIT_PLAN_ITEM_ID) not set"}
    try:
        import ms_graph
    except Exception as exc:  # noqa: BLE001
        return None, {"source": "graph", "error": f"ms_graph unavailable: {exc}"}
    if not ms_graph.have_ms_creds():
        return None, {"source": "graph", "error": "no Graph credentials configured"}
    token = ms_graph._get_token()
    if not token:
        return None, {"source": "graph", "error": "could not obtain Graph token"}
    import requests
    hdr = {"Authorization": f"Bearer {token}"}
    info: dict = {"source": "graph"}

    # Operations writes a NEW workbook about three times a day rather than
    # overwriting one — "PLAN DE VIAJES 11 DE SEPTIEMBRE 2026 (1).xlsm" and so
    # on — so list the folder and take the most recently modified .xlsm.
    if not item:
        url = (f"https://graph.microsoft.com/v1.0/drives/{drive}/items/{folder}/children"
               "?$select=id,name,size,lastModifiedDateTime&$top=200")
        try:
            r = requests.get(url, headers=hdr, timeout=60)
        except requests.RequestException as exc:
            return None, {**info, "error": f"folder listing failed: {exc}"}
        if r.status_code != 200:
            return None, {**info, "status": r.status_code,
                          "error": r.text[:300] or "folder listing refused "
                                                   "(app needs Files.Read.All)"}
        kids = [c for c in (r.json().get("value") or [])
                if str(c.get("name", "")).lower().endswith((".xlsm", ".xlsx"))]
        if not kids:
            return None, {**info, "error": "no .xlsm in the Plan de Viajes folder"}
        newest = max(kids, key=lambda c: str(c.get("lastModifiedDateTime") or ""))
        item = newest["id"]
        info.update(file=newest.get("name"), modified=newest.get("lastModifiedDateTime"),
                    candidates=len(kids))

    url = f"https://graph.microsoft.com/v1.0/drives/{drive}/items/{item}/content"
    try:
        r = requests.get(url, headers=hdr, timeout=60, allow_redirects=True)
    except requests.RequestException as exc:
        return None, {**info, "error": f"request failed: {exc}"}
    if r.status_code != 200:
        return None, {**info, "status": r.status_code,
                      "error": r.text[:300] or "download refused (app needs Files.Read.All)"}
    return r.content, {**info, "bytes": len(r.content)}


def load_plan() -> tuple[list[Trip], dict[str, Unit], dict]:
    """Trips + fleet from wherever the workbook can be fetched."""
    data, info = fetch_workbook_bytes()
    if not data:
        return [], {}, info
    trips, fleet, parse_info = load_workbook_trips(data)
    info.update(parse_info)
    return trips, fleet, info


# ── real trucks for the consolidation engine ─────────────────────────────────
# The engine plans against a hypothetical 88 m³ trailer. SIT already runs trucks
# to these hubs every week with space left on them, so the useful question is not
# "how many trailers would this need" but "which truck that is ALREADY GOING has
# room". Everything below answers that, and advises only — it never books.

def truck_hub(load: dict) -> str:
    """The Mexican hub a planned SIT truck is heading for, from its lanes."""
    for lane in load.get("lanes") or []:
        dest = lane.split("→")[-1].strip()
        hub = hub_for_destination(dest)
        if hub is not Hub.UNKNOWN:
            return hub.value
    return Hub.UNKNOWN.value


def spare_by_hub(trucks: list[dict]) -> dict:
    """Empty space already heading to each hub inside the horizon.

    This is the number the team cannot get anywhere else today: "you have
    118 m³ of paid-for empty space going to Guadalajara this week."
    """
    out: dict[str, dict] = {}
    for t in trucks:
        hub = truck_hub(t)
        b = out.setdefault(hub, {"hub": hub, "spare_m3": 0.0, "trucks": 0, "units": []})
        b["spare_m3"] = round(b["spare_m3"] + t.get("spare_m3", 0.0), 2)
        b["trucks"] += 1
        if t.get("unit") and t["unit"] not in b["units"]:
            b["units"].append(t["unit"])
    return dict(sorted(out.items(), key=lambda kv: -kv[1]["spare_m3"]))


def offer_trucks(loads: list[dict], trucks: list[dict],
                 today: Optional[dt.date] = None, slack_days: int = 3) -> list[dict]:
    """For each suggested load, the real SIT trucks that could carry it.

    A truck is a candidate when it is going to the same hub and loads before the
    load's depart-by date (plus a little slack, because a truck a day or two late
    is a conversation, not a no). `fits` says whether the whole load fits in the
    spare space; a partial is still worth showing — splitting one shipment onto
    a truck that is already rolling beats booking a new trailer.
    """
    today = today or dt.date.today()
    out = []
    for ld in loads:
        hub = ld.get("hub")
        depart = parse_date(ld.get("depart_by"))
        cands = []
        for t in trucks:
            if truck_hub(t) != hub or hub == Hub.UNKNOWN.value:
                continue
            when = parse_date(t.get("date"))
            if not when or when < today:
                continue
            if depart and when > depart + dt.timedelta(days=slack_days):
                continue
            spare = float(t.get("spare_m3") or 0)
            if spare <= 0:
                continue
            cands.append({"unit": t.get("unit"), "date": t.get("date"),
                          "driver": t.get("driver", ""), "spare_m3": spare,
                          "fill_pct": t.get("fill_pct"),
                          "fits": spare >= float(ld.get("m3") or 0)})
        cands.sort(key=lambda c: (not c["fits"], c["date"], -c["spare_m3"]))
        out.append({**ld, "trucks": cands[:4]})
    return out
