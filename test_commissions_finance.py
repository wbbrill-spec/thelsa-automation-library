"""Reader for Finance's gross-margin report (sheet `Base` of Margen x Expediente).

The first workbook built here has the same quirks as the real September 2026
report: a blank row above the header, a grand-total line at the bottom, a
US-dollar equivalent on peso invoices, blank currencies, lots typed differently
from their file, a mistyped exchange rate, a name where a job number belongs,
and diplomatic work billed to private individuals.

The second half takes the awkward cases one at a time (found by an independent
review on 8 Oct 2026): credit notes in a later month, jobs that net to zero,
amounts typed as text, and every place a customer name could slip through.
"""
import datetime as dt

import openpyxl
import pytest

from commissions import engine as E
from commissions import finance as F

D = dt.datetime
HEAD = ["Empresa", "Tipo", "Expediente", "Expediente 2", "Fecha", "Serie", "Folio", "Razon Social",
        "Factura que Cancela", "Neto Dólar Americano", "Moneda", "Tipo de Cambio", "Venta Neta",
        "Costo Proveedores", "Servs TRS a TMS", None, "COMENTARIOS"]
USD, MXN, EUR = "Dólar Americano", "Peso Mexicano", "Euros"
ROWS = [
    # entity, type, job, file, date, series, folio, bill-to, cancels, usd-equiv, ccy, fx, sales, sup, ic, -, comment
    ["LLC", "CORP-PART", "110001", "110001", D(2026, 5, 4), "CORP", 10, "Acme Corp", None, 10000, USD, 17.5, 175000, None, None, None, None],
    ["TMS SA", "CORP-PART", "110001", "110001", D(2026, 6, 20), "PROV-MXN", 77, "Supplier X", None, None, None, None, None, 90000, None, None, None],
    ["TRS", "CORP-PART", "110001", "110001", D(2026, 7, 2), "MEX", 5, "TRS", None, None, None, None, None, None, 20000, None, None],
    # peso invoice: the "dollar" column is only an equivalent
    ["SA", "CORP-PART", "110002", "110002", D(2026, 5, 9), "PART", 11, "Jane Roe", None, 5714.29, MXN, 17.5, 100000, None, None, None, None],
    # cancel and rebill across two months, blank currency (dollars to the cent)
    [None, "CORP-PART", "110003", "110003", D(2026, 5, 15), "CORP", 12, "Big Co", None, 20000.00, None, None, 344000.1234, None, None, None, None],
    [None, "CORP-PART", "110003", "110003", D(2026, 6, 3), "NCCORP", 3, "Big Co", "CORP 12", -20000.00, None, None, -346000.5678, None, None, None, None],
    [None, "CORP-PART", "110003", "110003", D(2026, 6, 4), "CORP", 13, "Big Co", None, 24000.00, None, None, 415200.9876, None, None, None, None],
    ["TMS LLC", "CORP-PART", "110003", "110003", D(2026, 8, 1), "PROV-USD", 78, "Supplier Y", None, None, None, None, None, 250000, None, None, None],
    # blank currency, pesos to the cent
    [None, "CORP-PART", "110004", "110004", D(2026, 6, 6), "TMS", 1138, "Other Co", None, 8499.486483, None, None, 146977.40, None, None, None, None],
    # a file whose lots are typed differently
    ["LLC", "AGENTE", "110771", "110771", D(2026, 6, 1), "AGENTE", 20, "Agent GmbH", None, 7000, USD, 17.3, 121100, None, None, None, None],
    ["LLC", "AGENTE", "110771A", "110771", D(2026, 6, 2), "AGENTE", 21, "Agent GmbH", None, 2000, USD, 17.3, 34600, None, None, None, None],
    ["SA", "CORP-PART", "110771B", "110771", D(2026, 6, 9), "CORP-PART", 22, "John Doe", None, 5780.35, MXN, 17.3, 100000, None, None, None, None],
    # diplomatic work billed to an individual, and an embassy invoice typed in scope
    ["SA", "DIPLOM", "110930", "110930", D(2026, 8, 5), "DIPLOM", 30, "Mary Major", None, 3000, MXN, 17.0, 51000, None, None, None, None],
    ["SA", "CORP-PART", "110972", "110972", D(2026, 7, 30), "CORP", 1179, "EMBAJADA DE LOS ESTADOS UNIDOS DE AMERICA", None, 7571.43, MXN, 17.5, 132500, None, None, None, None],
    # a mistyped rate, a mistyped job number, a job with two types, cost with no sale
    ["LLC", "CORP-PART", "110633", "110633", D(2026, 5, 27), "CORP-PART", 122, "Co Z", None, 19300.82, USD, 12.6231, 243636.18, None, None, None, None],
    ["LLC", "AGENTE", "MARCEL", "MARCEL", D(2026, 5, 28), "AGENTE", 40, "Marcel Example", None, 100, USD, 17.5, 1750, None, None, None, None],
    ["LLC", "AGENTE", "110894", "110894", D(2026, 5, 29), "AGENTE", 41, "Agent Ltd", None, 1000, USD, 17.5, 17500, None, None, None, "SE LIBERO COMO AGENTE/ RECLASIFICACION A CORP/PART JUNIO"],
    ["TMS SA", "CORP-PART", "110894", "110894", D(2026, 6, 30), "PROV-MXN", 79, "Supplier Q", None, None, None, None, None, 5000, None, None, None],
    ["TMS SA", "CORP-PART", "110500", "110500", D(2026, 1, 12), "PROV-MXN", 80, "Supplier R", None, None, None, None, None, 44000, None, None, None],
]
# the grand total line, as Finance's sheet ends
ROWS.append([None] * 12 + [sum(r[12] or 0 for r in ROWS), sum(r[13] or 0 for r in ROWS),
                          sum(r[14] or 0 for r in ROWS), None, None])


def book(path, rows, sheet="Base"):
    wb = openpyxl.Workbook()
    wb.active.title = "Gross Margin per file"
    ws = wb.create_sheet(sheet)
    ws.append([None] * len(HEAD))
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def row(job, sales=None, usd=None, ccy=MXN, date=D(2026, 5, 4), typ="CORP-PART", series="CORP", folio=1,
        sup=None, ic=None, cancels=None, billto="Some Customer SA", file=None, fx=None, comment=None, entity="SA"):
    return [entity, typ, job, file if file is not None else str(job)[:6], date, series, folio, billto, cancels,
            usd, ccy, fx, sales, sup, ic, None, comment]


def cost(job, amount, date=D(2026, 7, 1), typ="CORP-PART", folio=900):
    return row(job, sup=amount, ccy=None, date=date, typ=typ, series="PROV-MXN", folio=folio, billto="A Supplier")


@pytest.fixture
def load(tmp_path):
    n = [0]

    def _load(rows):
        n[0] += 1
        return F.load_base(book(tmp_path / f"b{n[0]}.xlsx", rows))
    return _load


@pytest.fixture(scope="module")
def lines(tmp_path_factory):
    return F.load_base(book(tmp_path_factory.mktemp("fin") / "Margen x Expediente.xlsx", ROWS))


# The rule for the commission is dollars at the internal rate, 16.5, and most
# tests below use it. The other setting, dollars at the rate of the invoice
# day, has its own section at the end of the file.
S = E.Settings(sales_fx="internal")
SPOT = E.Settings(sales_fx="spot")


def by_job(cl):
    return {m: {ln.job: ln for ln in v} for m, v in cl.items()}


def kinds_of(problems):
    out = {}
    for x in problems:
        out.setdefault(x["kind"], []).append(x)
    return out


# ── the report as it arrives ─────────────────────────────────────────────────
def test_reads_documents_not_the_total_line(lines):
    assert len(lines) == len(ROWS) - 1
    assert sum(ln.net_sales_mxn for ln in lines) == pytest.approx(sum(r[12] or 0 for r in ROWS[:-1]))
    assert {ln.kind for ln in lines} == {"sale", "supplier_cost", "interco_cost"}
    assert lines[0].row == 3                              # traceable to the sheet


def test_wrong_layout_is_refused(tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Base"
    wb.active.append(["Expediente", "Venta Neta", "Something else"])
    wb.save(p)
    with pytest.raises(F.FinanceReportError):
        F.load_base(p)
    wb.active.title = "Hoja1"
    wb.save(p)
    with pytest.raises(F.FinanceReportError):
        F.load_base(p)
    with pytest.raises(F.FinanceReportError):             # a header and nothing under it
        F.load_base(book(tmp_path / "empty.xlsx", []))


def test_no_customer_name_is_kept(lines):
    blob = repr(lines)
    for name in ("Acme", "Jane Roe", "Big Co", "Mary Major", "EMBAJADA", "Marcel", "MARCEL", "Supplier",
                 "RECLASIFICACION"):
        assert name not in blob
    blob = (repr(F.problems(lines, S)) + repr(F.jobs(lines, S)) + repr(F.commission_lines(lines, S))).upper()
    assert "MARCEL" not in blob and "RECLASIF" not in blob


def test_type_belongs_to_the_lot_not_the_file(lines):
    j = F.jobs(lines, S)
    assert j[("110771", "AGENTE")]["type"] == "AGENTE"
    assert j[("110771A", "AGENTE")]["file"] == "110771"
    assert j[("110771B", "CORP-PART")]["billed_mxn"] == pytest.approx(100000)
    in_scope = {ln.job for m in F.commission_lines(lines, S).values() for ln in m}
    assert "110771B" in in_scope and "110771" not in in_scope and "110771A" not in in_scope


def test_dollars_use_the_internal_rate_pesos_do_not(lines):
    j = F.jobs(lines, S)
    a = j[("110001", "CORP-PART")]
    assert a["billed_fin_mxn"] == pytest.approx(175000)          # Finance: 10,000 x 17.5
    assert a["billed_mxn"] == pytest.approx(165000)              # commission: 10,000 x 16.5
    b = j[("110002", "CORP-PART")]                               # a peso invoice
    assert b["billed_mxn"] == b["billed_fin_mxn"] == pytest.approx(100000)


def test_blank_currency_is_worked_out_and_reported(lines):
    by = {(ln.job, ln.folio): ln for ln in lines}
    assert (by[("110003", "12")].currency, by[("110003", "12")].currency_source) == ("USD", "amounts")
    assert (by[("110004", "1138")].currency, by[("110004", "1138")].currency_source) == ("MXN", "amounts")
    assert by[("110001", "10")].currency_source == "stated"
    found = [p for p in F.problems(lines, S) if p["kind"] == "currency_missing"]
    assert {p["job"] for p in found} == {"110003", "110004"} and len(found) == 4


def test_cancel_and_rebill_nets_out_by_month(lines):
    cl = by_job(F.commission_lines(lines, S))
    may, jun = cl["2026-05"]["110003"], cl["2026-06"]["110003"]
    assert may.billed == pytest.approx(20000 * 16.5)
    assert jun.billed == pytest.approx((-20000 + 24000) * 16.5)
    # the job's cost follows its billing, so every line carries the job's margin %
    total_billed, total_cost = 24000 * 16.5, 250000
    for ln in (may, jun):
        assert ln.cost / ln.billed == pytest.approx(total_cost / total_billed)
    assert may.cost + jun.cost == pytest.approx(total_cost)


def test_margin_part_is_the_same_however_cost_is_spread(lines):
    s = E.Settings(budgets={"2026-05": 1e6, "2026-06": 1e6}, sales_fx="internal")
    cl = F.commission_lines(lines, s)
    parts = [E.calc_line(ln, 0.6, s).margin for m in ("2026-05", "2026-06")
             for ln in cl[m] if ln.job == "110003"]
    assert sum(parts) == pytest.approx((24000 * 16.5 - 250000) * 0.03)


def test_cost_arrives_later_than_the_invoice_and_still_counts(lines):
    may = by_job(F.commission_lines(lines, S))["2026-05"]
    assert may["110001"].cost == pytest.approx(110000)           # June supplier + July TRS
    assert may["110001"].cost_posted is True


def test_no_cost_posted_holds_the_margin_part(lines):
    jun = by_job(F.commission_lines(lines, S))["2026-06"]
    assert jun["110004"].cost_posted is False
    r = E.calc_line(jun["110004"], 0.6, S)
    assert r.margin_pending and r.margin == 0 and r.discipline > 0


def test_cost_with_nothing_billed_makes_no_line(lines):
    assert all(ln.job != "110500" for m in F.commission_lines(lines, S).values() for ln in m)


def test_problems_found(lines):
    kinds = kinds_of(F.problems(lines, S))
    assert [x["job"] for x in kinds["fx_outlier"]] == ["110633"]
    assert kinds["fx_outlier"][0]["understated_mxn"] > 90000
    assert [x["job"] for x in kinds["embassy_billed_in_scope"]] == ["110972"]
    assert [x["job"] for x in kinds["bad_job_number"]] == ["?ROW18"]
    assert [x["job"] for x in kinds["type_conflict"]] == ["110894"]
    assert {x["job"] for x in kinds["no_cost_posted"]} >= {"110002", "110004", "110771B", "110972", "110633"}
    assert {x["job"] for x in kinds["cost_without_sales"]} == {"110500", "110894"}
    for absent in ("unknown_type", "no_date", "credit_note_amount", "cost_not_counted",
                   "sale_and_cost_on_one_row", "no_dollar_amount", "dollar_amount_mismatch", "empty_line",
                   "duplicate_line", "negative_cost", "date_out_of_range", "credit_reference_unreadable"):
        assert absent not in kinds


def test_diplomatic_work_billed_to_a_person_stays_out(lines):
    """The embassy cannot be recognised by who was billed; the type decides."""
    assert all(ln.job != "110930" for m in F.commission_lines(lines, S).values() for ln in m)


def test_embassy_invoice_typed_in_scope_is_kept_but_marked(lines):
    jul = by_job(F.commission_lines(lines, S))["2026-07"]
    assert jul["110972"].note == "billed to the Embassy"


# ── credit notes and cost: no month may carry more cost than the job has ─────
def margin_parts(cl, job, s=S):
    return [E.calc_line(ln, 0.6, s).margin for m in sorted(cl) for ln in cl[m] if ln.job == job]


def test_credit_note_in_a_later_month_gets_no_cost_and_cost_adds_up(load):
    """+100,000 in May, -80,000 in June, cost 10,000. Spreading cost over the
    net (20,000) would put 50,000 of cost on May and -40,000 on June."""
    ls = load([row("110010", 100000, date=D(2026, 5, 4)),
               row("110010", -80000, date=D(2026, 6, 9), series="NCCORP", folio=2),
               cost("110010", 10000)])
    cl = by_job(F.commission_lines(ls, S))
    may, jun = cl["2026-05"]["110010"], cl["2026-06"]["110010"]
    assert (may.billed, may.cost) == (pytest.approx(100000), pytest.approx(10000))
    assert (jun.billed, jun.cost) == (pytest.approx(-80000), 0.0)
    assert 0 <= may.cost <= 10000
    assert sum(margin_parts(F.commission_lines(ls, S), "110010")) == pytest.approx((20000 - 10000) * 0.03)
    assert not kinds_of(F.problems(ls, S)).get("cost_not_counted")


def test_three_months_with_a_credit_in_the_middle(load):
    ls = load([row("110011", 60000, date=D(2026, 5, 4)),
               row("110011", -90000, date=D(2026, 6, 9), series="NCCORP", folio=2),
               row("110011", 140000, date=D(2026, 7, 9), folio=3),
               cost("110011", 50000, date=D(2026, 8, 1))])
    cl = by_job(F.commission_lines(ls, S))
    costs = [cl[m]["110011"].cost for m in ("2026-05", "2026-06", "2026-07")]
    assert costs == [pytest.approx(15000), 0.0, pytest.approx(35000)]
    assert sum(costs) == pytest.approx(50000)
    assert sum(margin_parts(F.commission_lines(ls, S), "110011")) == pytest.approx((110000 - 50000) * 0.03)


def test_job_that_nets_to_zero_pays_nothing_and_its_cost_is_reported(load):
    ls = load([row("110012", 50000, date=D(2026, 5, 4)),
               row("110012", -50000, date=D(2026, 6, 9), series="NCCORP", folio=2),
               cost("110012", 7000)])
    cl = F.commission_lines(ls, S)
    assert all(ln.cost == 0 and ln.cost_posted for m in cl.values() for ln in m)
    s = E.Settings(budgets={"2026-05": 1e6, "2026-06": 1e6})
    total = sum(E.calc_month(m, cl[m], 1e6, s).total for m in cl)
    assert total == pytest.approx(0, abs=1e-6)             # what May paid, June takes back
    found = kinds_of(F.problems(ls, S))["cost_not_counted"]
    assert [(x["job"], x["amount_mxn"]) for x in found] == [("110012", 7000.0)]
    assert "nets to zero" in found[0]["detail"]


def test_credit_note_alone_takes_back_all_three_parts(load):
    """The invoice was in an earlier report. Its margin part was paid then, so
    the credit note must take a margin part back even with no cost in sight."""
    ls = load([row("110013", -30000, date=D(2026, 6, 9), series="NCCORP", folio=2)])
    ln = F.commission_lines(ls, S)["2026-06"][0]
    assert ln.cost_posted is True
    r = E.calc_line(ln, 0.6, S)
    assert not r.margin_pending
    assert (r.invoicing, r.margin, r.discipline) == (pytest.approx(-540), pytest.approx(-900), pytest.approx(-90))
    with_cost = load([row("110013", -30000, date=D(2026, 6, 9), series="NCCORP", folio=2), cost("110013", 2000)])
    assert F.commission_lines(with_cost, S)["2026-06"][0].cost == 0
    found = kinds_of(F.problems(with_cost, S))["cost_not_counted"]
    assert "credit notes only" in found[0]["detail"]


def test_invoice_and_credit_note_of_a_job_with_no_cost_wait_together(load):
    ls = load([row("110014", 100000, date=D(2026, 5, 4)),
               row("110014", -40000, date=D(2026, 6, 9), series="NCCORP", folio=2)])
    cl = F.commission_lines(ls, S)
    res = [E.calc_line(ln, 0.6, S) for m in sorted(cl) for ln in cl[m]]
    assert all(r.margin_pending and r.margin == 0 for r in res)
    assert sum(r.discipline for r in res) == pytest.approx(60000 * 0.03 * 0.10)
    found = kinds_of(F.problems(ls, S))["no_cost_posted"]
    assert [(x["job"], x["amount_mxn"]) for x in found] == [("110014", 60000.0)]


def test_cost_on_the_same_row_as_a_sale_is_counted_and_flagged(load):
    ls = load([row("110015", 100000, sup=30000, ic=5000)])
    j = F.jobs(ls, S)[("110015", "CORP-PART")]
    assert (j["billed_mxn"], j["cost_mxn"]) == (pytest.approx(100000), pytest.approx(35000))
    assert F.commission_lines(ls, S)["2026-05"][0].cost == pytest.approx(35000)
    assert [x["job"] for x in kinds_of(F.problems(ls, S))["sale_and_cost_on_one_row"]] == ["110015"]


# ── amounts typed as text ────────────────────────────────────────────────────
def test_amounts_typed_as_text_are_read(load):
    ls = load([row("110020", "$100,000.00", usd="5,714.29", fx="17.50"),
               row("110020", "(17,500.00)", usd="(1,000.00)", date=D(2026, 6, 1), series="NCCORP", folio=2),
               row("110020", "2,500.00-", date=D(2026, 6, 2), series="NCCORP", folio=3),
               row("110020", sup=" 40,000 ", ccy=None, series="PROV-MXN", folio=4),
               row("110021", "-", ic="-", ccy=None, sup=10)])
    assert [ln.net_sales_mxn for ln in ls] == [100000.0, -17500.0, -2500.0, 0.0, 0.0]
    assert ls[0].usd_equiv == pytest.approx(5714.29) and ls[0].fx == 17.5
    assert ls[3].supplier_cost_mxn == 40000.0


@pytest.mark.parametrize("col,bad", [(12, "pendiente"), (12, "17,500.00 aprox"), (13, "ver nota"),
                                     (14, "N/A"), (9, "USD 100"), (11, "diecisiete")])
def test_an_amount_that_cannot_be_read_stops_the_load(load, col, bad):
    r = row("110022", 1000, usd=57.14, fx=17.5)
    r[col] = bad
    with pytest.raises(F.FinanceReportError) as e:
        load([r])
    assert "row 3" in str(e.value) and bad not in str(e.value)


# ── names must not get through ───────────────────────────────────────────────
def test_names_do_not_get_through_any_column(load):
    ls = load([
        row("12 MONKEYS RELOCATION", 1000, file="12 MONKEYS RELOCATION", billto="Twelve Monkeys"),
        row("3D MUDANZAS", 2000, file="3D", folio=2, billto="3D Mudanzas"),
        row("110030", 3000, file="PEREZ FAMILY", folio=3, comment="Llamar a Sr. Perez", billto="Perez"),
        row("110031", -500, folio=4, series="NCCORP", cancels="CORP 12 de Juan Perez", billto="Perez"),
        row("110032", 700, folio=5, typ="JUAN PEREZ", billto="Perez"),
        row("110033", 800, folio="Perez-22 bis", series="SR PEREZ", billto="Perez", entity="Perez SA"),
    ])
    everything = (repr(ls) + repr(F.jobs(ls, S)) + repr(F.problems(ls, S)) + repr(F.commission_lines(ls, S))).upper()
    for name in ("MONKEY", "MUDANZAS", "3D", "PEREZ", "JUAN", "LLAMAR", "TWELVE"):
        assert name not in everything
    assert [ln.job for ln in ls][:2] == ["?ROW3", "?ROW4"]
    assert ls[2].file == "110030" and ls[2].has_comment is True
    assert ls[3].cancels == ""
    assert ls[4].type == "?" and ls[5].series == "?" and ls[5].folio == "?" and ls[5].entity == "OTHER"
    kinds = kinds_of(F.problems(ls, S))
    assert len(kinds["bad_job_number"]) == 2 and [x["job"] for x in kinds["unknown_type"]] == ["110032"]
    # the money on those lines is still there to be seen
    assert sum(ln.net_sales_mxn for ln in ls) == pytest.approx(7000)


def test_a_credit_note_reference_is_kept_as_series_and_number_only(load):
    ls = load([row("110034", -500, series="NCCORP", folio=9, cancels="corp-125"),
               row("110034", -500, series="NCCORP", folio=10, cancels="CORP PART 7"),
               row("110034", -500, series="NCCORP", folio=11, cancels="TMS1128"),
               row("110034", -500, series="NCCORP", folio=12, cancels="CORP 0125"),
               row("110034", -500, series="NCCORP", folio=13, cancels="CORP-PART-PART361"),
               row("110034", -500, series="NCCORP", folio=14, cancels="CORP-PART 22")])
    assert [x.cancels for x in ls] == ["CORP 125", "", "TMS 1128", "CORP 125", "", "CORP-PART 22"]
    assert [x.cancels_unread for x in ls] == [False, True, False, False, True, False]
    unread = kinds_of(F.problems(ls, S))["credit_reference_unreadable"]
    assert len(unread) == 2 and "PART" not in repr(unread)


def test_document_numbers_are_compared_without_leading_zeros(load):
    ls = load([row("110035", 175000, usd=10000, ccy=USD, fx=17.5, folio="0125"),
               row("110035", -218750, usd=-12500, ccy=USD, fx=17.5, date=D(2026, 6, 2), series="NCCORP", folio=6,
                   cancels="CORP 125")])
    assert ls[0].folio == "125"
    assert [x["job"] for x in kinds_of(F.problems(ls, S))["credit_note_amount"]] == ["110035"]


# ── currency ─────────────────────────────────────────────────────────────────
def test_euro_invoice_goes_through_the_dollar_equivalent(load):
    ls = load([row("110040", 187259.34, usd=10700.53, ccy=EUR, fx=17.5)])
    assert ls[0].currency == "EUR"
    assert F.jobs(ls, S)[("110040", "CORP-PART")]["billed_mxn"] == pytest.approx(10700.53 * 16.5)


def test_blank_currency_follows_the_other_line_of_the_same_invoice(load):
    """Real report, job 110611: a second line of a euro invoice with no currency
    and round figures on both sides. Guessing pesos was 93,629.67 out."""
    ls = load([row("110041", 175000.00, usd=10000.00, ccy=EUR, fx=17.5, folio=125),
               row("110041", 87500.00, usd=5000.00, ccy=None, folio=125)])
    assert (ls[1].currency, ls[1].currency_source) == ("EUR", "same invoice")
    assert F.jobs(ls, S)[("110041", "CORP-PART")]["billed_mxn"] == pytest.approx(15000 * 16.5)
    found = kinds_of(F.problems(ls, S))["currency_missing"]
    assert len(found) == 1 and "same invoice" in found[0]["detail"]


def test_blank_currency_that_cannot_be_worked_out_uses_finance_pesos_and_says_so(load):
    ls = load([row("110042", 87500.00, usd=5000.00, ccy=None, folio=7)])       # both sides to the cent
    assert (ls[0].currency, ls[0].currency_source) == ("", "not stated")
    assert F.jobs(ls, S)[("110042", "CORP-PART")]["billed_mxn"] == pytest.approx(87500)
    found = kinds_of(F.problems(ls, S))["currency_missing"]
    assert "could not be worked out" in found[0]["detail"] and found[0]["amount_mxn"] == 87500.0


def test_currency_labels_are_read_whatever_the_spelling(load):
    labels = ["Dólar Americano", "dolar americano", "USD", " Dólares ", "Peso Mexicano", "pesos", "MXN", "Euros", "EUR"]
    ls = load([row("110043", 1000, usd=57.14, ccy=c, folio=i) for i, c in enumerate(labels)])
    assert [ln.currency for ln in ls] == ["USD"] * 4 + ["MXN"] * 3 + ["EUR"] * 2
    assert all(ln.currency_source == "stated" for ln in ls)


def test_unknown_currency_and_missing_dollar_amount_are_flagged(load):
    ls = load([row("110044", 20000, usd=900, ccy="Libra Esterlina"),
               row("110045", 17500, usd=None, ccy=USD, fx=17.5)])
    j = F.jobs(ls, S)
    assert j[("110044", "CORP-PART")]["billed_mxn"] == pytest.approx(20000)
    assert j[("110045", "CORP-PART")]["billed_mxn"] == pytest.approx(17500)
    kinds = kinds_of(F.problems(ls, S))
    assert [x["job"] for x in kinds["currency_missing"]] == ["110044"]
    assert "LIBRA" not in repr(kinds).upper()
    assert [x["job"] for x in kinds["no_dollar_amount"]] == ["110045"]


# ── credit notes that do not match what they cancel ──────────────────────────
def test_credit_note_whose_amount_is_the_job_number(load):
    """Real report, job 110611: the credit note's dollar amount is -110,611.00."""
    ls = load([row("110611", 1657017.60, usd=94686.72, ccy=USD, fx=17.5, folio=125),
               row("110611", -1935692.50, usd=-110611.00, ccy=USD, fx=17.5, date=D(2026, 6, 2),
                   series="NCCORP", folio=16, cancels="CORP 125"),
               cost("110611", 900000)])
    found = kinds_of(F.problems(ls, S))["credit_note_amount"]
    assert len(found) == 1 and found[0]["job"] == "110611" and "job number itself" in found[0]["detail"]


def test_credit_note_bigger_than_the_invoice_it_cancels(load):
    ls = load([row("110050", 175000, usd=10000, ccy=USD, fx=17.5, folio=5),
               row("110050", -218750, usd=-12500, ccy=USD, fx=17.5, date=D(2026, 6, 2),
                   series="NCCORP", folio=6, cancels="CORP 5")])
    found = kinds_of(F.problems(ls, S))["credit_note_amount"]
    assert [x["job"] for x in found] == ["110050"] and "12,500.00 USD" in found[0]["detail"]


def test_credit_notes_that_are_fine_are_left_alone(load):
    ls = load([
        # for part of the invoice (a discount)
        row("110051", 26716.52, usd=1524.55, ccy=USD, fx=17.5242, series="PART", folio=396),
        row("110051", -8715.60, usd=-500, ccy=USD, fx=17.4312, date=D(2026, 6, 2), series="NCPART", folio=61,
            cancels="PART 396"),
        # the full amount at another day's rate
        row("110052", 175000, usd=10000, ccy=USD, fx=17.5, folio=7),
        row("110052", -176000, usd=-10000, ccy=USD, fx=17.6, date=D(2026, 6, 2), series="NCCORP", folio=8,
            cancels="CORP 7"),
        # its invoice is in an earlier report
        row("110053", -5000, usd=-285, ccy=MXN, date=D(2026, 6, 2), series="NCCORP", folio=9, cancels="CORP 1"),
        # names the wrong invoice of the job, but the money matches the other one (real job 109964)
        row("110054", 10799.60, usd=617.73, ccy=MXN, fx=17.4828, series="TMS", folio=1128),
        row("110054", 159526.98, usd=9139.59, ccy=MXN, fx=17.4545, series="TMS", folio=1134),
        row("110054", -159526.98, usd=-9164.53, ccy=MXN, fx=17.407, date=D(2026, 6, 9), series="NCTMS", folio=1028,
            cancels="TMS 1128"),
        # a peso credit note whose dollar equivalent happens to be whole
        row("110055", -1925000, usd=-110000, ccy=MXN, fx=17.5, date=D(2026, 6, 9), series="NCCORP", folio=30),
    ])
    assert "credit_note_amount" not in kinds_of(F.problems(ls, S))


def test_one_invoice_split_over_two_jobs_shares_its_currency(load):
    """Real report: CORP 129 is a euro invoice on job 110612 and, with the
    currency left blank, on job 110611. Finance's pesos were 93,629.67 too high."""
    ls = load([row("110612", 912434.0611906705, usd=52100.386066959996, ccy=EUR, fx=17.513, folio=129,
                   date=D(2026, 7, 27)),
               row("110611", 1618693.4125202673, usd=92428.10555132, ccy=None, folio=129, date=D(2026, 7, 27)),
               row("110613", 500.123, usd=28.5771, ccy=None, folio=129, date=D(2026, 8, 3))])   # another day: not it
    assert (ls[1].currency, ls[1].currency_source) == ("EUR", "same invoice")
    assert F.jobs(ls, S)[("110611", "CORP-PART")]["billed_mxn"] == pytest.approx(92428.10555132 * 16.5)
    assert (ls[2].currency, ls[2].currency_source) == ("", "not stated")


def test_same_invoice_number_in_two_currencies_settles_nothing(load):
    ls = load([row("110056", 175000, usd=10000, ccy=USD, fx=17.5, folio=5),
               row("110057", 175000, usd=10000, ccy=MXN, fx=17.5, folio=5),
               row("110058", 87500.5512, usd=5000.0315, ccy=None, folio=5)])
    assert ls[2].currency_source == "not stated"


# ── dates ────────────────────────────────────────────────────────────────────
def test_a_sale_with_no_date_is_in_no_month_and_is_reported(load):
    ls = load([row("110060", 100000, date=None), row("110060", 50000, date=D(2026, 5, 4), folio=2),
               row("110061", 9000, date="mayo 2026", folio=3)])
    j = F.jobs(ls, S)[("110060", "CORP-PART")]
    assert j["billed_mxn"] == pytest.approx(150000) and j["undated_billed_mxn"] == pytest.approx(100000)
    cl = F.commission_lines(ls, S)
    assert list(cl) == ["2026-05"] and cl["2026-05"][0].billed == pytest.approx(50000)
    found = kinds_of(F.problems(ls, S))["no_date"]
    assert [(x["job"], x["amount_mxn"]) for x in found] == [("110060", 100000.0), ("110061", 9000.0)]


# ── from the workbook to the commission ──────────────────────────────────────
def test_end_to_end_the_month_adds_up(lines):
    s = E.Settings(budgets=dict(E.BUDGET_2026))
    cl = F.commission_lines(lines, s)
    assert sorted(cl) == ["2026-05", "2026-06", "2026-07"]
    months = {m: E.calc_month(m, cl[m], s.budgets[m], s) for m in cl}       # bookings = budget
    assert all(r.reach == pytest.approx(0.6) for r in months.values())
    billed = sum(r.billed for r in months.values())
    in_scope = [j for j in F.jobs(lines, s).values() if j["type"] == "CORP-PART"]
    assert billed == pytest.approx(sum(j["billed_mxn"] for j in in_scope))
    # invoicing 60 % and discipline 10 % of 3 % on everything billed
    assert sum(x.invoicing + x.discipline for r in months.values() for x in r.lines) == pytest.approx(billed * 0.03 * 0.7)
    # margin part = 3 % of gross margin, on the jobs whose cost is in the books
    with_cost = [j for j in in_scope if j["cost_mxn"] and j["billed_mxn"]]
    assert {j["job"] for j in with_cost} == {"110001", "110003"}
    assert sum(x.margin for r in months.values() for x in r.lines) == pytest.approx(
        sum(j["billed_mxn"] - j["cost_mxn"] for j in with_cost) * 0.03)
    # nothing is payable until collection is known
    assert all(r.payable_total == 0 for r in months.values())
    waiting = {x.line.job for r in months.values() for x in r.lines if x.margin_pending}
    assert waiting == {"110002", "110004", "110771B", "110972", "110633"}


# ── rows that are not documents ──────────────────────────────────────────────
def total(rows):
    return [None] * 12 + [sum(r[12] or 0 for r in rows), sum(r[13] or 0 for r in rows),
                          sum(r[14] or 0 for r in rows), None, None]


def test_blank_rows_do_not_end_the_sheet_and_totals_are_checked(load):
    a, b = row("110070", 100000), cost("110070", 30000)
    blank = [None] * len(HEAD)
    ls = load([a, blank, blank, total([a]), b, total([a, b])])       # a subtotal, then the grand total
    assert [(x.net_sales_mxn, x.cost_mxn) for x in ls] == [(100000.0, 0.0), (0.0, 30000.0)]
    labelled = total([a, b])
    labelled[7] = "TOTAL GENERAL"                                    # a label in the name column
    assert len(load([a, b, labelled])) == 2


def test_money_on_a_row_with_no_job_stops_the_load(load):
    """A credit note and a cost under an invoice, with the job and type cells
    left blank (or merged): reading on would give a job with no cost."""
    a = row("110071", 100000)
    stray_credit = row(None, -40000, typ=None, series="NCCORP", folio=2)
    stray_credit[3] = None
    stray_cost = row(None, sup=55000, typ=None, ccy=None, folio=3)
    stray_cost[3] = None
    for rows in ([a, stray_credit], [a, stray_cost], [a, stray_credit, stray_cost]):
        with pytest.raises(F.FinanceReportError) as e:
            load(rows)
        assert "no type and no job number" in str(e.value)
    wrong_total = total([a])
    wrong_total[12] += 500
    with pytest.raises(F.FinanceReportError):
        load([a, wrong_total])


def test_a_row_with_a_job_but_no_type_or_a_type_but_no_job_is_kept_and_reported(load):
    ls = load([row("110072", 1000, typ=None), row(None, 2000, folio=2)])
    assert [(x.job, x.type) for x in ls] == [("110072", "?"), ("?ROW4", "CORP-PART")]
    kinds = kinds_of(F.problems(ls, S))
    assert [x["job"] for x in kinds["unknown_type"]] == ["110072"]
    assert [x["job"] for x in kinds["bad_job_number"]] == ["?ROW4"]
    assert all(ln.job != "110072" for m in F.commission_lines(ls, S).values() for ln in m)


def test_job_numbers_are_read_however_the_cell_holds_them(load):
    ls = load([row(110771.0, 1000), row(" 110771b ", 2000, folio=2), row("110 771 B", 3000, folio=3),
               row(110772, 4000, folio=4.0)])
    assert [x.job for x in ls] == ["110771", "110771B", "110771B", "110772"]
    assert ls[1].file == "110771" and ls[3].folio == "4"
    assert F.jobs(ls, S)[("110771B", "CORP-PART")]["billed_mxn"] == pytest.approx(5000)


def test_invoices_and_credit_notes_are_counted(load):
    ls = load([row("110073", 1000), row("110073", 2000, folio=2), row("110073", -500, series="NCCORP", folio=3)])
    j = F.jobs(ls, S)[("110073", "CORP-PART")]
    assert (j["invoices"], j["credit_notes"]) == (2, 1)


def test_a_cost_with_no_date_still_counts(load):
    ls = load([row("110074", 100000), cost("110074", 30000, date=None)])
    assert F.commission_lines(ls, S)["2026-05"][0].cost == pytest.approx(30000)
    assert [x["job"] for x in kinds_of(F.problems(ls, S))["no_date"]] == ["110074"]


# ── cents left over by rounding ──────────────────────────────────────────────
@pytest.mark.parametrize("inv,cn,ccy", [((10055.801724137931, 181016.49799655174), (-10055.80, -181016.47), USD),
                                        ((None, 100000.01), (None, -100000.00), MXN)])
def test_a_cancelled_job_with_cents_left_over_does_not_carry_its_cost(load, inv, cn, ccy):
    """Real report, job 110563 rows 3-4: invoice and credit note differ by a
    fraction of a cent in dollars. Without this the whole cost of the job would
    sit on a line of three centavos and take 1,500 pesos of commission."""
    ls = load([row("110075", inv[1], usd=inv[0], ccy=ccy, fx=18.0012 if ccy == USD else None),
               row("110075", cn[1], usd=cn[0], ccy=ccy, date=D(2026, 5, 20), series="NCCORP", folio=2),
               cost("110075", 50000)])
    cl = F.commission_lines(ls, S)
    assert all(ln.cost == 0 for m in cl.values() for ln in m)
    assert sum(abs(E.calc_line(ln, 0.6, S).total) for m in cl.values() for ln in m) < 0.01
    found = kinds_of(F.problems(ls, S))["cost_not_counted"]
    assert found[0]["amount_mxn"] == 50000.0 and "nets to zero" in found[0]["detail"]


def test_same_month_cancellation_is_called_what_it_is(load):
    ls = load([row("110076", 50000), row("110076", -50000, series="NCCORP", folio=2), cost("110076", 900)])
    kinds = kinds_of(F.problems(ls, S))
    assert "cost_without_sales" not in kinds and "nets to zero" in kinds["cost_not_counted"][0]["detail"]


# ── the dollar column drives the commission, so it is checked ────────────────
def test_a_dollar_amount_that_cannot_be_used_falls_back_to_pesos_and_is_reported(load):
    ls = load([row("110080", 175000, usd=0, ccy=USD, fx=17.5),
               row("110081", -175000, usd=10000, ccy=USD, fx=17.5, series="NCCORP", folio=2),
               row("110082", 175000, usd=None, ccy=EUR, folio=3)])
    j = F.jobs(ls, S)
    assert [j[(k, "CORP-PART")]["billed_mxn"] for k in ("110080", "110081", "110082")] == [175000, -175000, 175000]
    found = kinds_of(F.problems(ls, S))["no_dollar_amount"]
    assert [x["job"] for x in found] == ["110080", "110081", "110082"]
    assert "zero" in found[0]["detail"] and "opposite sign" in found[1]["detail"]


def test_dollars_and_pesos_that_do_not_agree_are_reported(load):
    ls = load([row("110083", 175000, usd=1000, ccy=USD, fx=17.5),          # a zero dropped from the dollars
               row("110084", 175000, usd=10000, ccy=USD, fx=17.5, folio=2),
               row("110085", 243636.18, usd=19300.82, ccy=USD, fx=12.6231, folio=3),   # rate typed wrong, amounts agree
               row("110086", 900000, usd=10000, ccy=USD, folio=4)])                    # no rate on the line: 90 a dollar
    found = kinds_of(F.problems(ls, S))["dollar_amount_mismatch"]
    assert [x["job"] for x in found] == ["110083", "110086"]
    assert F.jobs(ls, S)[("110083", "CORP-PART")]["billed_mxn"] == pytest.approx(16500)


def test_a_line_with_no_pesos_is_reported(load):
    ls = load([row("110087", None, usd=10000, ccy=USD, fx=17.5), row("110088", 1000, folio=2)])
    assert ls[0].kind == "empty"
    found = kinds_of(F.problems(ls, S))["empty_line"]
    assert [x["job"] for x in found] == ["110087"] and "10,000.00" in found[0]["detail"]


def test_the_same_document_twice_is_reported(load):
    ls = load([row("110089", 1000), row("110089", 1000), cost("110089", 7300, folio=55), cost("110089", 7300, folio=55),
               cost("110089", 7300, folio=56)])
    found = kinds_of(F.problems(ls, S))["duplicate_line"]
    assert [x["detail"][:12] for x in found] == ["rows 3, 4 ar", "rows 5, 6 ar"]
    assert F.jobs(ls, S)[("110089", "CORP-PART")]["cost_mxn"] == pytest.approx(21900)     # still counted


def test_negative_cost_is_counted_and_reported(load):
    ls = load([row("110090", 100000), cost("110090", -14000)])
    assert F.commission_lines(ls, S)["2026-05"][0].cost == pytest.approx(-14000)
    found = kinds_of(F.problems(ls, S))["negative_cost"]
    assert [(x["job"], x["amount_mxn"]) for x in found] == [("110090", -14000.0)]


def test_a_date_in_the_wrong_year_is_reported(load):
    ls = load([row("110091", 1000), row("110091", 2000, folio=2), row("110091", 3000, folio=3, date=D(2062, 5, 4)),
               cost("110091", 10, date=D(1900, 1, 1))])
    found = kinds_of(F.problems(ls, S))["date_out_of_range"]
    assert sorted(x["detail"][:22] for x in found) == ["row 5: dated 2062-05-0", "row 6: dated 1900-01-0"]


# ── amounts written in a way that can be read two ways ───────────────────────
@pytest.mark.parametrize("bad", ["1.234,56", "1234,56", "100.000", "nan", "inf", "1e5", "12,34.00", "#N/A", True])
def test_ambiguous_or_impossible_amounts_stop_the_load(load, bad):
    with pytest.raises(F.FinanceReportError) as e:
        load([row("110092", bad)])
    assert "row 3" in str(e.value)


def test_a_cell_that_holds_no_number_at_all():
    """A workbook cannot store "not a number"; such a cell, like a formula that
    was never calculated, arrives blank and the line is reported as empty."""
    for bad in (float("nan"), float("inf"), -float("inf")):
        with pytest.raises(F.FinanceReportError):
            F._num(bad, 7, "Venta Neta")


def test_an_exchange_rate_may_have_three_decimals(load):
    ls = load([row("110093", "17,513.00", usd="1,000.00", ccy=USD, fx="17.513")])
    assert ls[0].fx == 17.513 and ls[0].net_sales_mxn == 17513.0


# ── single words that could be a name ────────────────────────────────────────
def test_a_single_word_does_not_get_through_either(load):
    ls = load([row("3M", 1000, billto="x"), row("7UP", 2000, folio=2), row("24 HR", 3000, folio=3),
               row("110094", 4000, series="GARCIA", folio="LOPEZ"),
               row("110094", -100, series="NCCORP", folio=5, cancels="GARCIA 12"),
               row("110095", 5000, series="C000", folio=6, entity="GARCIA")])
    everything = (repr(ls) + repr(F.jobs(ls, S)) + repr(F.problems(ls, S)) + repr(F.commission_lines(ls, S))).upper()
    for word in ("3M", "7UP", "24HR", "GARCIA", "LOPEZ", "C000"):
        assert word not in everything
    assert [x.job for x in ls[:3]] == ["?ROW3", "?ROW4", "?ROW5"]
    assert (ls[3].series, ls[3].folio, ls[4].cancels, ls[5].series, ls[5].entity) == ("?", "?", "", "?", "OTHER")


def test_error_messages_carry_no_text_from_the_workbook(tmp_path):
    p = book(tmp_path / "s.xlsx", [row("110096", 1000)], sheet="Clientes Garcia")
    with pytest.raises(F.FinanceReportError) as e:
        F.load_base(p)
    assert "Garcia" not in str(e.value)
    with pytest.raises(E.SettingsError) as e:
        E.to_mxn(100, "Libra de Garcia", S, usd_equivalent=5)
    assert "Garcia" not in str(e.value)


def test_the_commission_uses_16_5_unless_told_otherwise(lines):
    """Bill, 9 Oct 2026: the day's rate would raise the pesos payable. 1 dollar
    = 16.5 pesos stays the rule; "spot" has to be asked for."""
    assert E.Settings().sales_fx == "internal" and E.Settings().usd_mxn == 16.5
    assert F.jobs(lines)[("110001", "CORP-PART")]["billed_mxn"] == pytest.approx(165000)
    assert F.jobs(lines, E.Settings())[("110001", "CORP-PART")]["billed_mxn"] == pytest.approx(165000)
    assert "currency_missing" in kinds_of(F.problems(lines))


# ── the other setting: dollars at the rate of the invoice day ────────────────
def test_with_spot_a_dollar_sale_counts_at_finances_peso_amount(lines):
    j = F.jobs(lines, SPOT)
    a = j[("110001", "CORP-PART")]
    assert a["billed_mxn"] == a["billed_fin_mxn"] == pytest.approx(175000)      # 10,000 dollars at 17.5
    for job in j.values():
        assert job["billed_mxn"] == pytest.approx(job["billed_fin_mxn"])
    may = by_job(F.commission_lines(lines, SPOT))["2026-05"]["110001"]
    assert (may.billed, may.cost) == (pytest.approx(175000), pytest.approx(110000))
    r = E.calc_line(may, 0.6, SPOT)
    assert r.margin == pytest.approx((175000 - 110000) * 0.03)                  # the margin Finance sees


def test_cancel_and_rebill_at_the_days_rates(lines):
    cl = by_job(F.commission_lines(lines, SPOT))
    assert cl["2026-05"]["110003"].billed == pytest.approx(344000.1234)
    assert cl["2026-06"]["110003"].billed == pytest.approx(-346000.5678 + 415200.9876)


def test_with_spot_the_currency_of_a_sale_raises_no_point(lines, load):
    kinds = kinds_of(F.problems(lines, SPOT))
    assert "currency_missing" not in kinds and "no_dollar_amount" not in kinds
    # the same sheet at the internal rate does raise them
    assert "currency_missing" in kinds_of(F.problems(lines, S))
    ls = load([row("110100", 20000, usd=900, ccy="Libra Esterlina"), row("110101", 17500, usd=None, ccy=USD, fx=17.5)])
    assert F.problems(ls, SPOT) == [p for p in F.problems(ls, SPOT) if p["kind"] == "no_cost_posted"]


def test_with_spot_a_mistyped_rate_does_change_the_commission_and_says_so(lines):
    found = kinds_of(F.problems(lines, SPOT))["fx_outlier"]
    assert [x["job"] for x in found] == ["110633"] and "too low" in found[0]["detail"]
    assert "too low" not in kinds_of(F.problems(lines, S))["fx_outlier"][0]["detail"]
    j = F.jobs(lines, SPOT)[("110633", "CORP-PART")]
    assert j["billed_mxn"] == pytest.approx(243636.18)                          # Finance's (understated) pesos


def test_with_spot_dollars_and_pesos_that_disagree_are_still_reported(load):
    ls = load([row("110102", 175000, usd=1000, ccy=USD, fx=17.5),               # a zero dropped from the dollars
               row("110103", 175000, usd=0, ccy=USD, fx=17.5, folio=2),
               row("110104", -175000, usd=10000, ccy=USD, fx=17.5, series="NCCORP", folio=3),
               row("110105", 175000, usd=10000, ccy=USD, fx=17.5, folio=4)])
    found = kinds_of(F.problems(ls, SPOT))["dollar_amount_mismatch"]
    assert [x["job"] for x in found] == ["110102", "110103", "110104"]
    assert all("uses the pesos" in x["detail"] for x in found)
    j = F.jobs(ls, SPOT)
    assert [j[(k, "CORP-PART")]["billed_mxn"] for k in ("110102", "110103", "110104")] == [175000, 175000, -175000]


def test_every_other_check_still_runs_on_a_dollar_line_with_spot(load):
    """A dollar line with an unusable dollar amount must not skip its other checks."""
    ls = load([row("110106", -5000, usd=None, ccy=USD, series="NCCORP", folio=2, cancels="see the email",
                   billto="EMBAJADA DE LOS ESTADOS UNIDOS", date=None)])
    kinds = kinds_of(F.problems(ls, SPOT))
    assert {"credit_reference_unreadable", "embassy_billed_in_scope", "no_date"} <= set(kinds)


def test_the_setting_must_be_one_of_the_two():
    with pytest.raises(E.SettingsError):
        E.Settings(sales_fx="budget").check()
    assert E.Settings(sales_fx="internal").check().sales_fx == "internal"


def test_with_spot_a_cancelled_dollar_job_does_not_carry_its_cost(load):
    """Invoice 10,000 dollars at 17.50, credit note at 17.60: in pesos the job
    is -1,000, not zero. It is still a cancelled job, so its cost stays off the
    lines and is reported; only the exchange difference is left."""
    ls = load([row("110110", 175000, usd=10000, ccy=USD, fx=17.5),
               row("110110", -176000, usd=-10000, ccy=USD, fx=17.6, date=D(2026, 6, 2), series="NCCORP", folio=2,
                   cancels="CORP 1"),
               cost("110110", 50000),
               # the other way round: the difference is positive, and it is not "billed, no cost yet"
               row("110111", 176000, usd=10000, ccy=USD, fx=17.6, folio=3),
               row("110111", -175000, usd=-10000, ccy=USD, fx=17.5, date=D(2026, 6, 2), series="NCCORP", folio=4,
                   cancels="CORP 3")])
    cl = F.commission_lines(ls, SPOT)
    assert all(ln.cost == 0 and ln.cost_posted for m in cl.values() for ln in m)
    total = sum(E.calc_line(ln, 0.6, SPOT).total for m in cl.values() for ln in m)
    assert abs(total) < 1e-6                                   # -1,000 and +1,000 of exchange difference
    kinds = kinds_of(F.problems(ls, SPOT))
    assert [(x["job"], x["amount_mxn"]) for x in kinds["cost_not_counted"]] == [("110110", 50000.0)]
    assert "no_cost_posted" not in kinds
    # a job that is partly credited is not cancelled: its cost counts
    part = load([row("110112", 175000, usd=10000, ccy=USD, fx=17.5),
                 row("110112", -88000, usd=-5000, ccy=USD, fx=17.6, date=D(2026, 6, 2), series="NCCORP", folio=2),
                 cost("110112", 50000)])
    assert sum(ln.cost for m in F.commission_lines(part, SPOT).values() for ln in m) == pytest.approx(50000)
