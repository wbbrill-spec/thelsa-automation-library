"""Which invoices are collected, read from Finance's report of customer payments.

The fixture has the shape of CONTPAQ's "Análisis de IVA de pagos de clientes"
and carries made-up customer names and bank references, so the tests can check
that none of it is ever kept or shown."""
import datetime as dt
import io
import json

import openpyxl
import pytest

from commissions import engine as E
from commissions import finance as F
from commissions import payments as P
from commissions import report as R
from test_commissions_finance import D, MXN, USD, cost, load, row  # noqa: F401  (load is a fixture)

NAMES = ("ACME HIDDEN SA", "REP Banorte MN 1078037800", "CR416", "Jane Secret")
S = E.Settings(budgets=dict(E.BUDGET_2026))
HEAD = ["Fecha", "Serie", "Folio", "Código Cliente", "Cliente", "Concepto", "Importe pagado", None, "Serie", "Folio", "Total"]


def pay(date, folio, amount, series="TMS", total=None):
    return [date, "REPBBVA", 7959, "CR416", "ACME HIDDEN SA", "REP Banorte MN 1078037800", amount, f"{series}{folio}",
            series, folio, amount if total is None else total]


def sheet(ws, currency, rows, company="THELSA RELOCATION SOLUTIONS", period="Del: 01/ENE/2026 al: 30/SEP/2026", pad=3,
          total=None):
    ws.append(["CONTPAQ i"] + [None] * pad + [company])
    ws.append([f"Moneda: {currency}"] + [None] * pad + ["ANÁLISIS DE IVA DE PAGOS DE CLIENTES (CRONOLÓGICAMENTE)"])
    ws.append([None] * (pad + 1) + [period])
    ws.append(["Impuesto a considerar: I.V.A."])
    ws.append([None, "PAGO", None, None, None, None, None, None, "FACTURA"])
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    num = lambda v: v if isinstance(v, (int, float)) else 0
    ws.append([None] * 6 + [sum(num(r[6]) for r in rows) if total is None else total, None, None, None,
                            sum(num(r[10]) for r in rows)])


def report(tmp_path, mxp=(), usd=(), name="pay.xlsx", **kw):
    wb = openpyxl.Workbook()
    wb.active.title = "USD"
    sheet(wb.active, "Dólar Americano", list(usd), pad=2, **kw)
    sheet(wb.create_sheet("MXP"), "Peso Mexicano", list(mxp), **kw)
    p = tmp_path / name
    wb.save(p)
    return p


def test_payments_report_is_read_without_customers_or_bank_references(tmp_path):
    p = report(tmp_path, mxp=[pay(46086, 1138, 725082.24), pay(dt.datetime(2026, 4, 1), 1139, 352781.53),
                              pay(46112, 1139, 444.45, total=999.0)], usd=[pay(46044, 1120, 22685.05)])
    got, info = P.load(p)
    assert info == {"company": "TRS", "from": "2026-01-01", "to": "2026-09-30", "currencies": ["USD", "MXN"],
                    "paid": {"USD": 22685.05, "MXN": 1078308.22}}
    assert [(x.date, x.currency, x.series, x.folio, x.paid) for x in got] == [
        ("2026-01-22", "USD", "TMS", "1120", 22685.05), ("2026-03-05", "MXN", "TMS", "1138", 725082.24),
        ("2026-04-01", "MXN", "TMS", "1139", 352781.53), ("2026-03-31", "MXN", "TMS", "1139", 444.45)]
    kept = json.dumps(P.to_records(got, info))
    for name in NAMES:
        assert name not in kept
    assert P.from_records(json.loads(kept)) == (got, info)
    c = P.collected(got)
    assert c[("TMS", "1139")] == {"MXN": pytest.approx(353225.98), "USD": 0.0, "last": "2026-04-01", "n": 2}


def test_a_payments_report_that_cannot_be_read_is_refused_without_quoting_it(tmp_path):
    cases = {
        "the amount paid is not a number": dict(mxp=[pay(46086, 1138, "a lot")]),
        "the invoice series or folio cannot be read": dict(mxp=[pay(46086, "Jane Secret", 10.0)]),
        "but the sheet's own total is": dict(mxp=[pay(46086, 1138, 10.0)], total=11.0),
        "the currency of the report is not US dollars or Mexican pesos": None,
        "does not look like CONTPAQ's report of customer payments": "other",
    }
    for words, kw in cases.items():
        if kw is None:
            wb = openpyxl.Workbook()
            sheet(wb.active, "Euro", [pay(46086, 1138, 10.0)])
            p = tmp_path / "eur.xlsx"
            wb.save(p)
        elif kw == "other":
            wb = openpyxl.Workbook()
            wb.active.append(["Month", "File"])
            p = tmp_path / "other.xlsx"
            wb.save(p)
        else:
            p = report(tmp_path, **kw)
        with pytest.raises(P.PaymentReportError) as e:
            P.load(p)
        assert words in str(e.value) and e.value.es != str(e.value), words
        for name in NAMES:
            assert name not in str(e.value) and name not in e.value.es
    with pytest.raises(P.PaymentReportError):
        P.load(io.BytesIO(b"not a workbook"))
    for doc in ({}, {"payments": [{"date": "x"}], "info": {}}, {"payments": "no", "info": None}):
        with pytest.raises(P.PaymentReportError):
            P.from_records(doc)


def test_a_company_the_report_does_not_know_is_not_named(tmp_path):
    got, info = P.load(report(tmp_path, mxp=[pay(46086, 1138, 10.0)], company="Jane Secret Holdings"))
    assert info["company"] == "" and "Jane" not in json.dumps(P.to_records(got, info))
    with pytest.raises(P.PaymentReportError) as e:                         # two companies in one workbook
        wb = openpyxl.Workbook()
        sheet(wb.active, "Peso Mexicano", [pay(46086, 1138, 10.0)])
        sheet(wb.create_sheet("b"), "Dólar Americano", [pay(46086, 1139, 10.0)], company="THELSA INTERNATIONAL MOVING")
        wb.save(tmp_path / "two.xlsx")
        P.load(tmp_path / "two.xlsx")
    assert "different companies" in str(e.value)


# ── from payments to the status of each line ─────────────────────────────────
def _fin(load):
    return load([
        row("110001", sales=100000.0, date=D(2026, 6, 5), series="TMS", folio=1),            # collected in full (with IVA)
        row("110002", sales=50000.0, date=D(2026, 6, 9), series="TMS", folio=2),             # collected in part
        row("110003", sales=20000.0, date=D(2026, 6, 20), series="TMS", folio=3),            # covered, not collected
        row("110004", sales=30000.0, date=D(2026, 6, 21), series="PART", folio=9),           # a series no report covers
        row("110005", sales=175000.0, usd=10000.0, ccy=USD, fx=17.5, date=D(2026, 7, 2), series="TMS", folio=4),   # paid in dollars
        row("110006", sales=40000.0, date=D(2026, 7, 3), series="TMS", folio=5),             # two invoices in the month,
        row("110006", sales=10000.0, date=D(2026, 7, 28), series="TMS", folio=6),            # only one collected
        row("110007", sales=60000.0, date=D(2026, 7, 4), series="TMS", folio=7),             # cancelled and rebilled
        row("110007", sales=-60000.0, date=D(2026, 7, 6), series="NCTMS", folio=1, cancels="TMS 7"),
        row("110007", sales=61000.0, date=D(2026, 7, 8), series="TMS", folio=8),
        row("110008", sales=5000.0, date=D(2026, 10, 2), series="TMS", folio=10),            # after the report's period
        row("110009", sales=9000.0, date=D(2026, 6, 2), series="TMS", folio=11, typ="AGENTE"),
    ])


PAID = [pay(46180, 1, 116000.0), pay(46181, 2, 20000.0), pay(46210, 4, 11600.0), pay(46211, 5, 46400.0),
        pay(46212, 8, 61000.0), pay(46213, 11, 10440.0), pay(46214, 999, 5.0)]


def test_a_line_is_paid_when_every_invoice_behind_it_is_collected_in_full(load, tmp_path):
    lines = _fin(load)
    ps, info = P.load(report(tmp_path, mxp=[p for p in PAID if p[9] != 4], usd=[p for p in PAID if p[9] == 4]))
    status, problems, summary = R.auto_payments(lines, ps, R.payment_reach([(ps, info)]))
    assert {k: v["paid"] for k, v in status.items()} == {
        ("2026-06", "110001"): True, ("2026-06", "110002"): False, ("2026-06", "110003"): False,
        ("2026-07", "110005"): True, ("2026-07", "110006"): False, ("2026-07", "110007"): True}
    assert ("2026-06", "110004") not in status and ("2026-10", "110008") not in status      # not covered: no status
    assert status[("2026-06", "110001")]["date"] == "2026-06-07" and status[("2026-06", "110002")]["date"] == ""
    assert [(p["kind"], p["job"]) for p in problems] == [("partly_collected", "110002")] and "40%" in problems[0]["detail"]
    assert summary["invoices"] == 9 and summary["full"] == 4 and summary["part"] == 1 and summary["open"] == 2
    assert summary["covered"] == 7 and summary["series"] == ["TMS"]
    assert summary["not_in_report"] == ["TMS 11", "TMS 999"]      # an agent invoice and one Finance's report does not have


def test_report_pays_collected_lines_and_a_mark_on_the_page_wins(load, tmp_path):
    lines = _fin(load)
    ps, info = P.load(report(tmp_path, mxp=[p for p in PAID if p[9] != 4], usd=[p for p in PAID if p[9] == 4]))
    reach = R.payment_reach([(ps, info)])
    assert reach == {"TMS": "2026-09-30"}
    rep = R.build(lines, S, {}, collected=ps, collected_reach=reach)
    by = {(r["month"], r["job"]): r for r in rep["lines"]}
    assert [(by[k]["payment"], by[k]["payment_source"]) for k in (("2026-06", "110001"), ("2026-06", "110002"), ("2026-06", "110004"))] == [
        ("Paid", "file"), ("Pending", "file"), ("", "")]
    assert by[("2026-06", "110001")]["paid_on"] == "2026-06-07"
    paid = [r for r in rep["lines"] if r["payment"] == "Paid"]
    assert rep["total_all"]["payable"] == pytest.approx(sum(r["total"] for r in paid)) and len(paid) == 3
    assert rep["collection"]["full"] == 4 and [c["kind"] for c in rep["checks"]][0] == "partly_collected"
    # by hand: 110004 marked Paid, 110001 put back to Pending
    marked = R.build(lines, S, {}, collected=ps, collected_reach=reach,
                     payments={"2026-06|110004": "Paid", "2026-06|110001": "Pending"})
    by = {(r["month"], r["job"]): r for r in marked["lines"]}
    assert (by[("2026-06", "110004")]["payment"], by[("2026-06", "110004")]["payment_source"]) == ("Paid", "marked")
    assert (by[("2026-06", "110001")]["payment"], by[("2026-06", "110001")]["paid_on"]) == ("Pending", "")
    assert by[("2026-07", "110005")]["payment_mark"] == "" and by[("2026-07", "110005")]["payment"] == "Paid"
    # without any payments report nothing changes
    plain = R.build(lines, S, {})
    assert plain["collection"] is None and plain["total_all"]["payable"] == 0
    assert json.dumps(marked)


def test_a_report_covers_only_its_own_period(load, tmp_path):
    lines = _fin(load)
    ps, info = P.load(report(tmp_path, mxp=[pay(46180, 1, 116000.0)], period="Del: 01/ENE/2026 al: 30/JUN/2026"))
    status, _p, summary = R.auto_payments(lines, ps, R.payment_reach([(ps, info)]))
    assert set(status) == {("2026-06", "110001"), ("2026-06", "110002"), ("2026-06", "110003")}    # July is beyond it
    assert summary["open"] == 2
