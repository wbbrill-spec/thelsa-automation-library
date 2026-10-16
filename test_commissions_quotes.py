"""The booked quote of each file, read from two Moveware reports.

The fixtures have the shape of the real exports (print-style, no column
headings) and carry made-up names, so the tests can check that no name is ever
kept or shown."""
import datetime as dt
import io
import json

import openpyxl
import pytest

from commissions import engine as E
from commissions import finance as F
from commissions import quotes as Q
from commissions import report as R
from commissions import template_sheet as T
from test_commissions_finance import D, MXN, USD, book, cost, load, row  # noqa: F401  (load is a fixture)

NAMES = ("Jane Secret", "Acme Hidden SA", "jane@secret.example", "Supplier Shadow", "Rue Cachee 5")
S = E.Settings(budgets=dict(E.BUDGET_2026))


def J(job, branch="MEM", ctype="Private", won=dt.date(2026, 6, 10), ex=100000.0, inc=None, typ="IMP", status="W"):
    return dict(job=job, branch=branch, ctype=ctype, won=won, ex=ex, inc=ex if inc is None else inc, typ=typ, status=status)


def removals_rows(jobs, total_files=None, total_ex=None):
    """Rows of a Removals Summary: Branch > Company > Customer > Customer Type > Date > job."""
    rows = []
    for branch in sorted({j["branch"] for j in jobs}):
        mine = [j for j in jobs if j["branch"] == branch]
        tot = lambda js: [len(js), 100 * len(js), 3 * len(js), None, None, sum(j["ex"] for j in js), sum(j["inc"] for j in js)]
        rows.append([f"Branch: {branch}"] + tot(mine))
        for j in mine:
            rows.append(["Company: Acme Hidden SA"] + tot([j]))
            rows.append(["Customer: Jane Secret"] + tot([j]))
            rows.append([f"Customer Type: {j['ctype']}"] + tot([j]))
            rows.append([f"Date: {j['won']:%d/%m/%y}" if isinstance(j["won"], dt.date) else f"Date: {j['won']}"] + tot([j]))
            rows.append([j["job"], j["typ"], "N", "N", "FCL 20", "Jane Secret/Mrs", "Rue Cachee 5", "Mexico City",
                         "Acme Hidden SA", "Acme Hidden SA", branch, "MXMEX", 100, 3, j["status"], dt.time(0), dt.time(0),
                         j["ex"], j["inc"], "jane@secret.example"])
    rows.append(["Grand Total: ", len(jobs) if total_files is None else total_files, 3 * len(jobs),
                 sum(j["ex"] for j in jobs) if total_ex is None else total_ex, sum(j["inc"] for j in jobs)])
    rows.append(["Printed by Jane Secret"])
    return rows


def O(no="01", status="W", lines=((1000.0, 1300.0, None, None),), after=(), total=None):
    """A cost option. lines / after = (quoted cost, sale price, actual cost, date)."""
    return dict(no=no, status=status, lines=list(lines), after=list(after), total=total)


def costing_rows(jobs):
    """Rows of a Move Costing Sheet. jobs = [(job, move type, client type, [options])]."""
    rows = [["Won From: 01/01/26, Won To: 09/10/26, Type: DIN,EXP,EXT,IMP"]]
    line = lambda c, s_, a, d: ["Origin Services", 1, c, c, None, None if s_ is None else s_ - (c or 0), s_, None if a is None else 1, a, d,
                                "A", "N", "", "Supplier Shadow"]
    for job, move, ctype, options in jobs:
        rows.append(["Removal Plan", ":", job, "Name", ":", "Jane Secret", "Details", ":", "Rue Cachee 5", "Type", ":", move,
                     "RP Status : W", "Est Move : ", f"Client Type: {ctype}", "Volume : 35.00", "Nett Pounds: 9,000.00"])
        for o in options:
            rows.append(["Cost Option", ":", o["no"], "Door to door for Jane Secret", "Cost Option Details:_x000D_", "Surveyor: ",
                         "Cost Status", ":", o["status"], "Service : FCL 20", "Cubic Meter : 35.00", "Feet : 1,237.00",
                         "Values Inclusive of Tax", "Nett Pounds : 8,034.00", "None", "None", "None"])
            rows.append(["ESTIMATED COSTS", "ACTUAL COSTS", "ACCRUALS AND INVOICING", "Description", "Qty", "Rate", "Value",
                         "MU%", "Margin", "Sell Price", "Qty", "Value", "Date", "Acc", "Inv", "Invoice", "Creditor"])
            for l in o["lines"]:
                rows.append(line(*l))
            rows.append(["Profit Margin", None, None, None, None, None, None, None, None, None, "N", "F", "", ""])
            c = sum(l[0] or 0 for l in o["lines"]) if o["total"] is None else o["total"]
            s_ = sum(l[1] or 0 for l in o["lines"])
            rows.append(["Option Total", None, None, c, 30.0, s_ - c, s_, None, None, None, "N", "F", "", ""])
            for l in o["after"]:
                rows.append(line(*l))
    return rows


def xl(tmp_path, rows, name="r.xlsx"):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    p = tmp_path / name
    wb.save(p)
    return p


# ── Removals Summary ─────────────────────────────────────────────────────────
def test_removals_summary_is_read_without_its_names(tmp_path):
    jobs = [J("110001", "MEU", "Company", dt.date(2026, 6, 10), 10000.0, 11600.0, "EXP"),
            J("110002", "MEM", "Private", dt.date(2026, 7, 2), 250000.0), J("110003A", "MEM", "Blank", typ="IMA"),
            J("110004", "QRU", "Diplomatic"), J("110005", "QRU", "Agent")]
    got, control = Q.load_removals(xl(tmp_path, removals_rows(jobs)))
    assert [(x.job, x.move_type, x.branch, x.customer_type, x.won, x.value_ex, x.value_inc) for x in got][:2] == [
        ("110002", "IMP", "MEM", "Private", "2026-07-02", 250000.0, 250000.0),
        ("110003A", "IMA", "MEM", "Blank", "2026-06-10", 100000.0, 100000.0)]
    assert {x.job: x.customer_type for x in got} == {"110001": "Company", "110002": "Private", "110003A": "Blank",
                                                     "110004": "Diplomatic", "110005": "Agent"}
    assert control == {"listed": 5, "listed_ex": 560000.0, "total_files": 5, "total_ex": 560000.0}
    kept = json.dumps(Q.removals_to_records(got, control))
    for name in NAMES:
        assert name not in kept
    assert Q.removals_from_records(json.loads(kept)) == (got, control)
    assert Q.problems(got, control) == []


def test_removals_summary_reports_its_own_totals_when_they_differ(tmp_path):
    jobs = [J("110001"), J("110002")]
    got, control = Q.load_removals(xl(tmp_path, removals_rows(jobs)[:-2] + [["Grand Total: ", 3, 9, 1.0, 1.0]]))
    assert control["total_files"] == 2                      # the branch rows are what is compared, not the last line
    rows = removals_rows(jobs)
    rows[0][1], rows[0][6] = 3, 350000.0                    # the branch says three jobs, two are listed
    got, control = Q.load_removals(xl(tmp_path, rows, "b.xlsx"))
    p = Q.problems(got, control)
    assert [x["kind"] for x in p] == ["quote_totals"] and "3 jobs and 350,000.00" in p[0]["detail"]
    assert "2 jobs and 200,000.00" in p[0]["detail"] and "expedientes" in p[0]["detail_es"]


@pytest.mark.parametrize("change, words", [
    (lambda r: r.__setitem__(5, r[5][:17] + ["about 5000"] + r[5][18:]), "row 6: the value excluding tax is not a number"),
    (lambda r: r.__setitem__(5, r[5][:17] + [-5.0] + r[5][18:]), "row 6: a negative value"),
    (lambda r: r.__setitem__(5, r[5][:10] + ["QRU"] + r[5][11:]), "row 6: the branch of the job does not match"),
    (lambda r: r.__setitem__(5, r[5][:1] + ["IM-PORT"] + r[5][2:]), "row 6: the move type is not a Moveware code"),
    (lambda r: r.__setitem__(5, r[5][:14] + ["Won"] + r[5][15:]), "row 6: the status is not a single letter"),
    (lambda r: r.insert(6, list(r[5])), "rows 6 and 7: job 110001 is listed twice"),
    (lambda r: r.__setitem__(0, ["Branch: Jane Secret Branch"] + r[0][1:]), "row 1: the branch is not a short code"),
    (lambda r: r.__delitem__(slice(0, 5)), "row 1: a job is listed before its branch and date groups"),
])
def test_a_removals_summary_that_cannot_be_read_is_refused_without_quoting_it(tmp_path, change, words):
    rows = removals_rows([J("110001")])
    change(rows)
    with pytest.raises(Q.QuoteReportError) as e:
        Q.load_removals(xl(tmp_path, rows))
    assert words in str(e.value) and e.value.es and e.value.es != str(e.value)
    for name in NAMES:
        assert name not in str(e.value) and name not in e.value.es


def test_a_cell_that_is_not_a_plain_number_never_becomes_an_amount():
    assert Q._num(None, 3, "x", "x") is None and Q._num("  ", 3, "x", "x") is None and Q._num(5, 3, "x", "x") == 5.0
    for bad in (float("nan"), float("inf"), True, "12", dt.date(2026, 1, 1)):
        with pytest.raises(Q.QuoteReportError) as e:
            Q._num(bad, 7, "quoted cost", "el costo cotizado")
        assert str(e.value) == "row 7: the quoted cost is not a number" and "fila 7" in e.value.es


def test_other_workbooks_are_not_taken_for_a_removals_summary(tmp_path):
    for rows in ([["Month", "File"], ["June", "110001"]], costing_rows([("110001", "EXP", "Private", [O()])]), []):
        with pytest.raises(Q.QuoteReportError) as e:
            Q.load_removals(xl(tmp_path, rows))
        assert "does not look like the Moveware Removals Summary" in str(e.value)
    with pytest.raises(Q.QuoteReportError) as e:
        Q.load_removals(io.BytesIO(b"PK\x03\x04 not a workbook"))
    assert "could not be opened" in str(e.value)


def test_a_date_or_customer_type_that_cannot_be_read_is_kept_out_and_listed(tmp_path):
    rows = removals_rows([J("110001", won="sometime in June", ctype="Jane Secret")])
    got, control = Q.load_removals(xl(tmp_path, rows))
    assert got[0].won == "" and got[0].customer_type == "?"
    assert [p["kind"] for p in Q.problems(got, control)] == ["quote_no_date"]
    assert "Jane" not in json.dumps(Q.removals_to_records(got, control))


# ── Move Costing Sheet ───────────────────────────────────────────────────────
def test_costing_sheet_adds_the_lines_below_the_option_total(tmp_path):
    opts = [O("01", "C", [(900.0, 1000.0, None, None)]),
            O("02", "W", [(1000.0, 1300.0, 950.0, dt.datetime(2026, 7, 1)), (500.0, 700.0, None, None)],
              after=[(200.0, 260.0, 210.0, dt.datetime(2026, 8, 15))]),
            O("03", "W", [(300.0, 400.0, None, None)]), O("04", "L", [(1.0, 2.0, None, None)])]
    got, period = Q.load_costing(xl(tmp_path, costing_rows([("110001", "EXP", "Private", opts), ("110002A", "IMP", "", [O()])])))
    assert period == {"won_from": "2026-01-01", "won_to": "2026-10-09"}
    a = got[0]
    assert (a.job, a.move_type, a.client_type, [o.status for o in a.options]) == ("110001", "EXP", "Private", ["C", "W", "W", "L"])
    w = a.options[1]
    assert (w.cost, w.sell, w.lines, w.loaded, w.actual, w.last_actual, w.total_off) == (
        1700.0, 2260.0, 3, 2, 1160.0, "2026-08-15", False)
    assert got[1].client_type == "Blank"
    kept = json.dumps(Q.costing_to_records(got, period))
    for name in NAMES:
        assert name not in kept
    assert Q.costing_from_records(json.loads(kept)) == (got, period)


def test_an_option_total_that_is_not_the_sum_of_its_lines_is_listed(tmp_path):
    got, _ = Q.load_costing(xl(tmp_path, costing_rows([("110001", "EXP", "Private", [O(total=999.0)])])))
    assert got[0].options[0].total_off and got[0].options[0].cost == 999.0
    p = Q.problems(costings=got)
    assert [x["kind"] for x in p] == ["quote_total_off"] and "option 01" in p[0]["detail"]


def test_a_costing_sheet_that_cannot_be_read_is_refused(tmp_path):
    base = costing_rows([("110001", "EXP", "Private", [O()])])
    cases = {
        "no job number": [r if r[0] != "Removal Plan" else r[:2] + ["Jane Secret"] + r[3:] for r in base],
        "listed twice": base + base[1:],
        "before any removal plan": [base[0]] + base[2:],
        "\"Option Total\" lines instead of one": [r for r in base if r[0] != "Option Total"],
        "is not a number": [r if r[0] != "Origin Services" else r[:3] + ["a lot"] + r[4:] for r in base],
        "not a single letter": [r if r[0] != "Cost Option" else r[:8] + ["Won"] + r[9:] for r in base],
        "does not look like the Moveware Move Costing Sheet": removals_rows([J("110001")]),
    }
    for words, rows in cases.items():
        with pytest.raises(Q.QuoteReportError) as e:
            Q.load_costing(xl(tmp_path, rows))
        assert words in str(e.value), words
        for name in NAMES:
            assert name not in str(e.value) and name not in e.value.es


def test_stored_records_that_were_tampered_with_are_refused():
    for doc in ({}, {"jobs": [{"job": "x"}], "control": {}}, {"jobs": "no", "control": None}):
        with pytest.raises(Q.QuoteReportError):
            Q.removals_from_records(doc)
        with pytest.raises(Q.QuoteReportError):
            Q.costing_from_records(doc)


# ── the booked quote of a job ────────────────────────────────────────────────
def _reports(tmp_path, jobs, costings):
    rem, control = Q.load_removals(xl(tmp_path, removals_rows(jobs), "rem.xlsx"))
    cst, _ = Q.load_costing(xl(tmp_path, costing_rows(costings), "cst.xlsx"))
    return rem, control, cst


def test_quote_is_the_value_excluding_iva_and_the_w_options(tmp_path):
    rem, control, cst = _reports(tmp_path,
        [J("110001", "MEU", "Company", ex=10000.0, inc=11600.0, typ="EXP"), J("110002", "MEM", "Private", ex=250000.0),
         J("110003", "MEM", "Blank", ex=5000.0, typ="IMA"), J("110004", "MEM", ex=0.0)],
        [("110001", "EXP", "Company", [O("01", "W", [(6000.0, 8000.0, None, None)], after=[(500.0, 700.0, None, None)]),
                                       O("02", "W", [(1000.0, 1300.0, None, None)]), O("03", "C", [(9e6, 9e6, None, None)])]),
         ("110002", "IMP", "Private", [O("01", "L")]), ("110004", "IMP", "Private", [O("01", "W", [(700.0, 900.0, None, None)])])])
    q = Q.book(rem, cst, S)
    a = q["110001"]                                           # dollars at the internal rate, W options added up
    assert a["sale"] == pytest.approx(10000 * 16.5) and a["cost"] == pytest.approx(7500 * 16.5)
    assert (a["currency"], a["with_tax"], a["w_options"], a["status"], a["won"]) == ("USD", True, 2, "W", "2026-06-10")
    assert q["110002"]["sale"] == 250000.0 and q["110002"]["cost"] is None        # no W option: no quoted cost
    assert q["110003"]["cost"] is None and not Q.in_scope_by_moveware(q["110003"])
    assert Q.in_scope_by_moveware(q["110001"]) and Q.in_scope_by_moveware(q["110002"])
    z = q["110004"]                                           # no value in the Removals Summary: the costing's price
    assert (z["sale"], z["sale_from"], z["with_tax"]) == (900.0, "costing", True)


def test_a_job_on_the_wrong_branch_is_read_in_its_real_currency(tmp_path):
    rem, control, cst = _reports(tmp_path,
        [J("110766", "MEU", "Company", ex=288463.92), J("110633", "MEU", "Private", ex=19300.82),
         J("110900", "MEU", "Private", ex=300000.0, inc=348000.0), J("110901", "MEU", "Private", ex=300000.0),
         J("110902", "XXZ", "Private", ex=5.0)], [("110766", "IMP", "Company", [O()])])
    assert Q.currency_of("110766", "MEU") == "MXN" and Q.currency_of("110633", "MEU") == "CAD"
    # 110900: the sale price including IVA is what Finance invoiced in pesos, so it is pesos; 110901 is not near
    pesos = Q.looks_like_pesos(rem, {"110900": 348010.0, "110901": 5_100_000.0, "110766": 288463.92, "110902": 5.0})
    assert pesos == {"110900"}
    q = Q.book(rem, cst, S, pesos=pesos)
    assert q["110766"]["sale"] == pytest.approx(288463.92) and q["110766"]["cost"] == pytest.approx(1000.0)
    assert q["110633"]["sale"] is None                        # Canadian dollars: no rate here, so no quote
    assert q["110900"]["sale"] == pytest.approx(300000.0) and q["110901"]["sale"] == pytest.approx(300000 * 16.5)
    assert q["110902"]["sale"] is None                        # a branch that says nothing about the currency
    kinds = {p["job"]: p["detail"] for p in Q.problems(rem, control, cst, pesos=pesos)}
    assert set(kinds) == {"110766", "110633", "110900", "110902"}
    assert "CAD" in kinds["110633"] and "invoiced in pesos" in kinds["110900"] and "does not end in U or M" in kinds["110902"]


# ── in the report ────────────────────────────────────────────────────────────
def _finance(load, extra=()):
    """Three files. 110001 is invoiced in June and again in July; 110002 in July; 110003 in July."""
    return load([row("110001", sales=100000.0, date=D(2026, 6, 5), folio=1), row("110001", sales=50000.0, date=D(2026, 7, 9), folio=2),
                 row("110002", sales=80000.0, date=D(2026, 7, 3), folio=3), row("110003", sales=20000.0, date=D(2026, 7, 20), folio=4),
                 cost("110001", 60000.0, date=D(2026, 7, 1))] + list(extra))


def test_a_files_quote_is_shown_and_counted_once_where_it_is_first_invoiced(load, tmp_path):
    lines = _finance(load)
    rem, control, cst = _reports(tmp_path, [J("110001", ex=200000.0), J("110002", ex=90000.0), J("110003", ex=30000.0)],
                                 [("110001", "IMP", "Private", [O("01", "W", [(120000.0, 232000.0, None, None)])])])
    quotes, qp = R.book_quotes(lines, rem, control, cst, S)
    rep = R.build(lines, S, {}, quotes=quotes, quote_problems=qp)
    jun, jul = rep["months"]
    assert (jun["bookings_source"], jun["bookings"], jun["quote"]["files"], jun["quote"]["found"]) == ("quotes", 200000.0, 1, 1)
    assert (jul["bookings_source"], jul["bookings"]) == ("quotes", 120000.0)       # 110001 is not counted again
    assert jun["reach"] == pytest.approx(200000.0 / S.budgets["2026-06"] * 0.6)
    assert rep["placeholder_months"] == [] and rep["quotes_loaded"]
    by = {(r["month"], r["job"]): r for r in rep["lines"]}
    first, again = by[("2026-06", "110001")], by[("2026-07", "110001")]
    assert (first["status"], first["sale_price"], first["quote_cost"]) == ("Ganado", 200000.0, 120000.0)
    assert (again["status"], again["sale_price"], again["quote_cost"], again["quote_month"]) == ("Ganado", None, None, "2026-06")
    assert by[("2026-07", "110002")]["quote_cost"] is None
    assert rep["total_all"]["total"] == pytest.approx(sum(r["total"] for r in rep["lines"]))
    # the same lines without the quotes: both months on the placeholder
    plain = R.build(lines, S, {})
    assert plain["placeholder_months"] == ["2026-06", "2026-07"] and not plain["quotes_loaded"]
    assert plain["months"][0]["quote"] is None and "status" in plain["lines"][0]


def test_a_month_with_a_file_that_has_no_quote_keeps_the_placeholder(load, tmp_path):
    lines = _finance(load)
    rem, control, cst = _reports(tmp_path, [J("110001", ex=200000.0), J("110002", ex=90000.0)], [("110001", "IMP", "Private", [O()])])
    quotes, qp = R.book_quotes(lines, rem, control, cst, S)
    rep = R.build(lines, S, {}, quotes=quotes, quote_problems=qp)
    jun, jul = rep["months"]
    assert jun["bookings_source"] == "quotes" and jul["bookings_source"] == "placeholder"
    assert jul["quote"]["missing"] == ["110003"] and jul["quote"]["bookings"] == 90000.0 and jul["bookings"] is None
    assert jul["reach"] == pytest.approx(0.6) and rep["placeholder_months"] == ["2026-07"]
    missing = [c for c in rep["checks"] if c["kind"] == "quote_missing"][0]
    assert [i["job"] for i in missing["items"]] == ["110003"] and missing["routine"]
    # a figure typed on the page wins over both
    typed = R.build(lines, S, {"2026-06": 1_000_000.0, "2026-07": 2_000_000.0}, quotes=quotes, quote_problems=qp)
    assert [m["bookings_source"] for m in typed["months"]] == ["entered", "entered"]
    assert typed["months"][0]["bookings"] == 1_000_000.0


def test_bookings_by_month_won_is_the_other_basis(load, tmp_path):
    lines = _finance(load)
    rem, control, cst = _reports(tmp_path,
        [J("110001", ex=200000.0, won=dt.date(2026, 6, 1)), J("110002", ex=90000.0, won=dt.date(2026, 6, 20)),
         J("110003", ex=30000.0, won=dt.date(2026, 7, 2), ctype="Blank"), J("110050", ex=7e6, won=dt.date(2026, 6, 2), typ="IMA"),
         J("110051", ex=8e6, won=dt.date(2026, 6, 3), ctype="Diplomatic"), J("110052", ex=40000.0, won=dt.date(2026, 7, 9), ctype="Company")],
        [("110001", "IMP", "Private", [O()])])
    s = E.Settings(budgets=dict(E.BUDGET_2026), bookings_basis="won")
    quotes, qp = R.book_quotes(lines, rem, control, cst, s)
    rep = R.build(lines, s, {}, quotes=quotes, quote_problems=qp)
    # agent and diplomatic jobs are not bookings; a blank customer type counts when Finance has the file in scope
    assert [m["bookings"] for m in rep["months"]] == [290000.0, 70000.0]
    assert rep["bookings_basis"] == "won" and [m["quote"]["won"] for m in rep["months"]] == [290000.0, 70000.0]
    with pytest.raises(E.SettingsError):
        E.Settings(bookings_basis="shipped").check()


def test_points_about_the_quotes_are_listed(load, tmp_path):
    lines = _finance(load)
    rem, control, cst = _reports(tmp_path,
        [J("110001", ex=200000.0, inc=232000.0), J("110002", ex=90000.0, typ="IMA"), J("110003", ex=30000.0, ctype="Diplomatic")],
        [("110001", "IMP", "Private", [O()])])
    quotes, qp = R.book_quotes(lines, rem, control, cst, S)
    rep = R.build(lines, S, {}, quotes=quotes, quote_problems=qp)
    kinds = {c["kind"]: [i["job"] for i in c["items"]] for c in rep["checks"]}
    assert kinds["moveware_type_differs"] == ["110002", "110003"] and kinds["quote_with_tax"] == ["110001"]
    assert kinds["quote_no_cost"] == ["110002", "110003"]
    assert all(c["label"] != c["kind"] and c["label_es"] != c["label"] for c in rep["checks"])
    assert json.dumps(rep)


# ── on the Excel sheet ───────────────────────────────────────────────────────
def test_sheet_shows_the_quote_and_sums_the_sale_prices(load, tmp_path):
    lines = _finance(load)
    rem, control, cst = _reports(tmp_path, [J("110001", ex=200000.0), J("110002", ex=90000.0)],
                                 [("110001", "IMP", "Private", [O("01", "W", [(120000.0, 232000.0, None, None)])])])
    quotes, qp = R.book_quotes(lines, rem, control, cst, S)
    rep = R.build(lines, S, {}, quotes=quotes, quote_problems=qp)
    wb = T.build_book(rep, S)
    jun, jul = wb["Jun 2026"], wb["Jul 2026"]
    assert [jun[f"{c}4"].value for c in "BCDE"] == ["110001", "Ganado", 200000.0, 120000.0]
    assert jun["D5"].value == "=SUM(D4:D4)" and jun["I5"].value == "=IF(H5=0,0,D5/H5)"       # bookings = the sale prices
    notes = lambda ws: " ".join(str(c.value) for r_ in ws.iter_rows(min_col=2, max_col=2) for c in r_ if c.value)
    assert "sum of the sale prices of the files first invoiced" in notes(jun) and "from Moveware" in notes(jun)
    # July: 110001 again (status only), 110002 with its price, 110003 with no quote, so the budget stands in
    rows = {jul[f"B{r}"].value: [jul[f"{c}{r}"].value for c in "CDE"] for r in (4, 5, 6)}
    assert rows == {"110001": ["Ganado", None, None], "110002": ["Ganado", 90000.0, None], "110003": [None, None, None]}
    assert jul["D7"].value == S.budgets["2026-07"] and "110003" in notes(jul) and "set equal to the budget" in notes(jul)
    text = " ".join(str(c.value) for ws in wb for r_ in ws.iter_rows() for c in r_ if c.value is not None)
    for name in NAMES:
        assert name not in text


def test_a_cost_line_entered_twice_on_the_valid_option_is_listed_without_its_text(tmp_path):
    jobs = [("110001", "EXP", "Company", [
                O("01", "W", lines=((1000.0, 1300.0, None, None), (1000.0, 1300.0, None, None), (500.0, 650.0, 480.0, None)),
                  after=((0.0, 0.0, 480.0, None),)),
                O("02", "C", lines=((700.0, 900.0, None, None), (700.0, 900.0, None, None)))]),   # cancelled option: ignored
            ("110002", "IMP", "Private", [O("01", "W", lines=((1000.0, 1300.0, None, None), (900.0, 1100.0, None, None)))])]
    got, _ = Q.load_costing(xl(tmp_path, costing_rows(jobs)))
    by = {j.job: j for j in got}
    assert [(d["kind"], d["amount"], d["times"]) for d in by["110001"].duplicates] == [("quoted", 1000.0, 2),
                                                                                      ("actual", 480.0, 2)]
    assert by["110002"].duplicates == ()
    back, _ = Q.costing_from_records(Q.costing_to_records(got, {"won_from": "", "won_to": ""}))
    assert back == got and "Supplier" not in str(Q.costing_to_records(got, {}))
    probs = [p for p in Q.problems(costings=got) if p["kind"] == "costing_duplicate"]
    assert [p["job"] for p in probs] == ["110001", "110001"] and "aparece 2 veces" in probs[0]["detail_es"]


def test_a_move_type_not_on_the_settings_list_is_listed(load, tmp_path):
    import dataclasses
    lines = _finance(load)
    jobs = [J("110001", typ="XUSA"), J("110002", typ="IMP")]
    rem, control, cst = _reports(tmp_path, jobs, [("110002", "IMP", "Private", [O()])])
    s = E.Settings(budgets=dict(E.BUDGET_2026))
    quotes, qp = R.book_quotes(lines, rem, control, cst, s)
    rep = R.build(lines, s, {}, quotes=quotes, quote_problems=qp)
    flagged = {i["job"] for c in rep["checks"] if c["kind"] == "move_type_unlisted" for i in c["items"]}
    assert "110001" in flagged and "110002" not in flagged
    s2 = dataclasses.replace(s, move_types=s.move_types + (("XUSA", "tms"),))
    quotes, qp = R.book_quotes(lines, rem, control, cst, s2)
    rep = R.build(lines, s2, {}, quotes=quotes, quote_problems=qp)
    assert not any(c["kind"] == "move_type_unlisted" for c in rep["checks"])
