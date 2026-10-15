"""Pay runs: Lupita submits, Mario approves or returns, Rogelio records payment;
nothing is paid twice, and what changes later goes in the next run."""
import dataclasses
import io

import openpyxl
import pytest

from commissions import engine as E
from commissions import finance as F
from commissions import report as R
from commissions import signoff as SO
from commissions import store
from commissions import web as W
from test_commissions_web import LUPITA, OUTSIDER, _lines_of, client, login, upload, xlsx  # noqa: F401

MARIO, ROGELIO, BILL = "mariotorres@thelsa.com", "rogeliofranco@thelsa.com", "bbrill@thelsa.com"
S = E.Settings(budgets=dict(E.BUDGET_2026))


def rep(*lines):
    """A report with only what the sign-off reads: month, file, line total, payment."""
    return {"lines": [{"month": m, "job": j, "total": t, "payment": p} for m, j, t, p in lines]}


def submit(runs, who, d, note=""):
    return SO.act(runs, "submit", who, S, current=d, expect=d["fingerprint"], note=note)


@pytest.fixture(autouse=True)
def _roles(monkeypatch):
    for k in ("COMMISSION_REVIEWERS", "COMMISSION_APPROVERS", "COMMISSION_FINANCE", "COMMISSION_ADMINS",
              "COMMISSION_TEST_SELF_APPROVAL"):
        monkeypatch.delenv(k, raising=False)


def test_draft_pays_collected_new_lines_and_splits_to_the_cent():
    d = SO.draft(rep(("2026-05", "110001", 500.0, "Paid"), ("2026-06", "110002", 1000.0, "Paid"),
                     ("2026-06", "110003", 700.0, "Pending"), ("2026-07", "110004", 333.33, "Paid")), [], S, "2026-06")
    assert [(x["month"], x["job"]) for x in d["lines"]] == [("2026-06", "110002"), ("2026-07", "110004")]
    assert d["due"] == pytest.approx(1333.33) and d["months"] == ["2026-06", "2026-07"]
    assert [p["name"] for p in d["people"]] == [n for n, _ in S.sales_people] + [n for n, _ in S.admin_people]
    assert sum(p["share"] for p in d["people"]) == pytest.approx(1.0)
    assert all(round(p["amount"], 2) == p["amount"] for p in d["people"])
    assert abs(SO.control(d)) <= 0.05                          # only rounding can leave cents over


def test_the_whole_chain_and_nothing_is_paid_twice():
    runs = []
    r1 = rep(("2026-06", "110002", 1000.0, "Paid"), ("2026-07", "110004", 200.0, "Pending"))
    d = SO.draft(r1, runs, S, "2026-06")
    run = SO.act(runs, "submit", LUPITA, S, current=d, expect=d["fingerprint"], note="June checked")
    assert run["state"] == "submitted" and run["log"][0]["by"] == LUPITA and not run["log"][0]["as_admin"]
    with pytest.raises(SO.SignOffError):                       # only one run waits at a time
        submit(runs, LUPITA, d)
    with pytest.raises(SO.SignOffError):                       # Lupita cannot approve her own run
        SO.act(runs, "approve", LUPITA, S, current=d, run_id=1)
    with pytest.raises(SO.SignOffError):                       # Rogelio cannot pay before approval
        SO.act(runs, "pay", ROGELIO, S, run_id=1)
    SO.act(runs, "approve", MARIO, S, current=SO.draft(r1, runs, S, "2026-06"), run_id=1)
    assert runs[0]["state"] == "approved" and "changed_at_approval" not in runs[0]
    assert SO.draft(r1, runs, S, "2026-06")["lines"] == []    # approved = paid out: nothing left
    SO.act(runs, "pay", ROGELIO, S, run_id=1, paid_on="2026-10-30")
    assert runs[0]["state"] == "paid" and runs[0]["paid_on"] == "2026-10-30"
    # later: July is collected and June's line grows (its cost arrived: margin part)
    r2 = rep(("2026-06", "110002", 1250.0, "Paid"), ("2026-07", "110004", 200.0, "Paid"))
    d2 = SO.draft(r2, runs, S, "2026-06")
    assert [(x["job"], x["due"], x["adjust"]) for x in d2["lines"]] == [("110002", 250.0, True), ("110004", 200.0, False)]
    assert d2["due"] == pytest.approx(450.0)


def test_return_and_withdraw_count_for_nothing_and_need_a_reason():
    runs = []
    r = rep(("2026-06", "110002", 1000.0, "Paid"))
    d = SO.draft(r, runs, S, "2026-06")
    submit(runs, LUPITA, d)
    with pytest.raises(SO.SignOffError):
        SO.act(runs, "return", MARIO, S, run_id=1, note="  ")
    SO.act(runs, "return", MARIO, S, run_id=1, note="check 110002 cost")
    assert runs[0]["state"] == "returned" and SO.draft(r, runs, S, "2026-06")["due"] == pytest.approx(1000.0)
    submit(runs, LUPITA, SO.draft(r, runs, S, "2026-06"))
    SO.act(runs, "withdraw", LUPITA, S, run_id=2)
    assert [x["state"] for x in runs] == ["returned", "withdrawn"] and runs[1]["id"] == 2


def test_figures_that_moved_are_refused_at_submit_and_noted_at_approval():
    runs = []
    d = SO.draft(rep(("2026-06", "110002", 1000.0, "Paid")), runs, S, "2026-06")
    moved = SO.draft(rep(("2026-06", "110002", 1100.0, "Paid")), runs, S, "2026-06")
    with pytest.raises(SO.SignOffError):
        SO.act(runs, "submit", LUPITA, S, current=moved, expect=d["fingerprint"])
    SO.act(runs, "submit", LUPITA, S, current=d, expect=d["fingerprint"])
    SO.act(runs, "approve", MARIO, S, current=moved, run_id=1)
    assert runs[0]["due"] == pytest.approx(1000.0)            # Mario approved what was submitted
    assert runs[0]["changed_at_approval"]["on_these_lines"] == pytest.approx(100.0)
    assert SO.draft(rep(("2026-06", "110002", 1100.0, "Paid")), runs, S, "2026-06")["due"] == pytest.approx(100.0)


def test_nothing_to_pay_or_a_negative_run_cannot_be_submitted_and_take_back_is_never_automatic():
    runs = []
    with pytest.raises(SO.SignOffError):
        submit(runs, LUPITA, SO.draft(rep(), runs, S, "2026-06"))
    with pytest.raises(SO.SignOffError):
        submit(runs, LUPITA, SO.draft(rep(("2026-06", "1", -50.0, "Paid")), runs, S, "2026-06"))
    submit(runs, LUPITA, SO.draft(rep(("2026-06", "1", 50.0, "Paid")), runs, S, "2026-06"))
    SO.act(runs, "approve", MARIO, S, run_id=1)
    d = SO.draft(rep(("2026-06", "1", 50.0, "Pending")), runs, S, "2026-06")
    assert d["lines"] == [] and d["warnings"][0]["job"] == "1"     # listed, not clawed back


def test_roles_come_from_render_and_an_admin_step_says_so(monkeypatch):
    assert SO.roles_of(LUPITA) == ["review"] and SO.roles_of(MARIO) == ["approve"] and SO.roles_of(ROGELIO) == ["finance"]
    assert SO.roles_of(BILL) == ["admin"] and not SO.may(OUTSIDER, "submit")
    monkeypatch.setenv("COMMISSION_APPROVERS", "other.boss@thelsa.com")
    assert not SO.may(MARIO, "approve") and SO.may("other.boss@thelsa.com", "approve")
    runs = []
    submit(runs, BILL, SO.draft(rep(("2026-06", "1", 50.0, "Paid")), runs, S, "2026-06"))
    assert runs[0]["log"][0]["as_admin"] is True
    with pytest.raises(SO.SignOffError):
        SO.act(runs, "cancel", BILL, S, run_id=1, note="x")        # only an approved run can be cancelled
    with pytest.raises(SO.SignOffError):
        SO.act(runs, "approve", BILL, S, run_id=1)                  # not even an admin approves his own run
    SO.act(runs, "approve", "other.boss@thelsa.com", S, run_id=1)
    SO.act(runs, "cancel", BILL, S, run_id=1, note="paid outside the app")
    assert runs[0]["state"] == "cancelled" and SO.paid_so_far(runs) == {}


def test_instruction_workbook_adds_up():
    runs = []
    submit(runs, LUPITA, SO.draft(rep(("2026-06", "110002", 1234.57, "Paid")), runs, S, "2026-06"))
    wb = SO.instruction_book(runs[0])
    ws = wb["Payment instruction"]
    amounts = [ws.cell(row=r, column=4).value for r in range(5, 5 + len(runs[0]["people"]))]
    assert sum(amounts) == pytest.approx(1234.57, abs=0.05) and wb["Lines"]["B2"].value == "110002"
    assert ws.cell(row=5 + len(amounts) + 1, column=4).value == pytest.approx(1234.57)


# ── through the page ─────────────────────────────────────────────────────────
def _mark_new_lines_paid(c):
    rep_ = _lines_of(c)
    new = [x for x in rep_["lines"] if x["month"] >= R.NEW_FROM]
    c.post("/commissions/payments", data={"csrf": "tok", **{f"p_{x['month']}_{x['job']}": "Paid" for x in new}})
    return _lines_of(c)


def test_page_flow_lupita_mario_rogelio(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    full = _mark_new_lines_paid(client)
    page = client.get("/commissions").data.decode()
    assert "Reviewed: send to Mario for approval" in page and "Approve run" not in page
    fp = W._draft(F.from_records(store.get("finance_report")[0]["records"]), W.settings(), [])["fingerprint"]
    assert client.post("/commissions/signoff", data={"csrf": "tok", "step": "approve", "run": "1"}).status_code == 302
    assert store.get("pay_runs")[0] is None                   # nothing to approve, nothing stored
    client.post("/commissions/signoff", data={"csrf": "tok", "step": "submit", "expect": fp, "note": "ok"})
    runs = store.get("pay_runs")[0]
    assert runs[0]["state"] == "submitted" and runs[0]["due"] == pytest.approx(full["total_new"]["payable"])
    login(client, MARIO)
    page = client.get("/commissions").data.decode()
    assert "Approve run 1" in page and "Reviewed: send to Mario" not in page
    client.post("/commissions/signoff", data={"csrf": "tok", "step": "approve", "run": "1"})
    login(client, ROGELIO)
    assert "Mark paid" in client.get("/commissions").data.decode()
    client.post("/commissions/signoff", data={"csrf": "tok", "step": "pay", "run": "1", "paid_on": "2026-10-31"})
    runs = store.get("pay_runs")[0]
    assert runs[0]["state"] == "paid" and [e["by"] for e in runs[0]["log"]] == [LUPITA, MARIO, ROGELIO]
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/commissions/signoff/1/instruction").data))
    assert wb.sheetnames == ["Payment instruction", "Lines"]
    assert client.get("/commissions/signoff/9/instruction").status_code == 404
    assert "Nothing new to pay" in client.get("/commissions").data.decode()


def test_signoff_needs_the_token_and_the_list(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    _mark_new_lines_paid(client)
    assert client.post("/commissions/signoff", data={"csrf": "bad", "step": "submit"}).status_code == 400
    login(client, OUTSIDER)
    assert client.post("/commissions/signoff", data={"csrf": "tok", "step": "submit"}).status_code == 404
    assert client.get("/commissions/signoff/1/instruction").status_code == 404


def test_the_same_person_cannot_submit_and_approve_or_approve_and_pay(monkeypatch):
    monkeypatch.setenv("COMMISSION_REVIEWERS", "both@thelsa.com")
    monkeypatch.setenv("COMMISSION_APPROVERS", "both@thelsa.com")
    monkeypatch.setenv("COMMISSION_FINANCE", "both@thelsa.com")
    runs = []
    submit(runs, "both@thelsa.com", SO.draft(rep(("2026-06", "1", 50.0, "Paid")), runs, S, "2026-06"))
    with pytest.raises(SO.SignOffError):
        SO.act(runs, "approve", "both@thelsa.com", S, run_id=1)
    SO.act(runs, "approve", BILL, S, run_id=1)
    with pytest.raises(SO.SignOffError):
        SO.act(runs, "pay", BILL, S, run_id=1)
    SO.act(runs, "pay", "both@thelsa.com", S, run_id=1)
    monkeypatch.setenv("COMMISSION_TEST_SELF_APPROVAL", "1")        # a test copy of the app only
    submit(runs, BILL, SO.draft(rep(("2026-06", "2", 50.0, "Paid")), runs, S, "2026-06"))
    SO.act(runs, "approve", BILL, S, run_id=2)
    assert runs[1]["state"] == "approved"


def test_a_settings_change_never_pays_again_on_a_line_already_paid():
    runs = []
    submit(runs, LUPITA, SO.draft(rep(("2026-06", "A", 300.0, "Paid")), runs, S, "2026-06"))
    SO.act(runs, "approve", MARIO, S, run_id=1)
    # rate goes from 3 % to 4 %: today's report says 400 for A; under run 1's settings it is still 300
    now = rep(("2026-06", "A", 400.0, "Paid"), ("2026-07", "B", 80.0, "Paid"))
    d = SO.draft(now, runs, S, "2026-06", totals_by_run={1: {("2026-06", "A"): 300.0}})
    assert [(x["job"], x["due"]) for x in d["lines"]] == [("B", 80.0)]
    # but a real change on A (its cost arrived) under run 1's settings still goes in the next run
    d = SO.draft(now, runs, S, "2026-06", totals_by_run={1: {("2026-06", "A"): 330.0}})
    assert [(x["job"], round(x["due"], 2)) for x in d["lines"]] == [("A", 30.0), ("B", 80.0)]


def test_a_draft_seen_with_other_settings_cannot_be_submitted():
    d = SO.draft(rep(("2026-06", "A", 300.0, "Paid")), [], S, "2026-06")
    other = dataclasses.replace(S, sales_share=0.5, admin_share=0.5)
    d2 = SO.draft(rep(("2026-06", "A", 300.0, "Paid")), [], other, "2026-06")
    assert d["fingerprint"] != d2["fingerprint"]
    with pytest.raises(SO.SignOffError):
        SO.act([], "submit", LUPITA, S, current=d2, expect=d["fingerprint"])
    with pytest.raises(SO.SignOffError):
        SO.act([], "submit", LUPITA, S, current=d)                  # no fingerprint sent: refused


def test_settings_change_through_the_page_does_not_repay(client, tmp_path):
    from test_commissions_settings import form
    login(client)
    upload(client, xlsx(tmp_path))
    full = _mark_new_lines_paid(client)
    fp = W._draft(F.from_records(store.get("finance_report")[0]["records"]), W.settings(), [])["fingerprint"]
    client.post("/commissions/signoff", data={"csrf": "tok", "step": "submit", "expect": fp})
    login(client, MARIO)
    client.post("/commissions/signoff", data={"csrf": "tok", "step": "approve", "run": "1"})
    login(client, BILL)
    client.post("/commissions/settings", data=form(rate="4"))
    assert W.settings().rate == pytest.approx(0.04)
    nxt = W._draft(F.from_records(store.get("finance_report")[0]["records"]), W.settings(), W._runs())
    assert nxt["lines"] == [] and full["total_new"]["payable"] > 0


def test_two_steps_at_once_cannot_undo_each_other(client):
    import threading
    store.put("pay_runs", [])
    errors = []

    def add(n):
        try:
            store.update("pay_runs", lambda rs: (rs + [n], None), default=[])
        except Exception as e:                      # noqa: BLE001
            errors.append(e)
    ts = [threading.Thread(target=add, args=(i,)) for i in range(20)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert not errors and sorted(store.get("pay_runs")[0]) == list(range(20))
