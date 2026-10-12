"""A month of commission in the layout of Finance's own workbook.

The formulas written to the sheet are the template's. Excel is not available
here to calculate them, so these tests check the sheet's structure and
formulas, and check that the engine gives the template's numbers for the same
inputs (the workbook itself is recalculated and compared cell by cell with the
template before it is sent to anyone)."""
import openpyxl
import pytest

from commissions import engine as E
from commissions import template_sheet as T

S = E.Settings(budgets=dict(E.BUDGET_2026))


def sheet(rows=T.SAMPLE_ROWS, budget=T.SAMPLE_BUDGET, **kw):
    ws = openpyxl.Workbook().active
    return ws, T.write(ws, rows, S, budget, **kw)


def test_same_blocks_and_columns_as_the_template():
    ws, info = sheet()
    assert (ws["A2"].value, ws["K2"].value, ws["Q1"].value, ws["W2"].value) == (
        "BOOKED QUOTE", "REAL", "GLOBAL COMMISSION", "COMMISSION DISTRIBUTION")
    assert ws["Q2"].value == 0.03
    assert [ws.cell(row=3, column=c).value for c in (1, 2, 4, 8, 11, 12, 16, 17, 19, 21)] == [
        "Month", "File", "Sale prices", "Budget per month", "Billed", "Actual Cost", "Payment status",
        "% Sales/Billed", "Gross Margin", "Total Commission"]
    assert [ws.cell(row=3, column=c).value for c in range(23, 28)] == T.DIST_HEAD
    assert info == {"first": 4, "last": 10, "total": 11, "targets": 13, "sales_total": 15, "adm_first": 16,
                    "adm_total": 22}                                  # the template's own row numbers
    assert (ws["D14"].value, ws["D15"].value, ws["D16"].value) == (0.6, 0.3, 0.1)
    assert ws["H11"].value == 500000 and ws["A11"].value == "Total Bookings"


def test_formulas_are_the_templates():
    ws, _ = sheet()
    assert ws["J4"].value == "=$I$11*$D$14"                           # sales reach
    assert ws["I11"].value == "=IF(H11=0,0,D11/H11)"                  # budget achievement
    assert ws["D11"].value == "=SUM(D4:D10)"
    assert ws["Q4"].value == "=(K4*$Q$2)*J4"                          # invoicing part
    assert ws["R4"].value == '=IF(O4="","",O4/$D$15)'                 # margin achievement
    assert ws["S4"].value == '=IF(OR(R4="",K4=0),0,((K4*$Q$2)*$D$15)*R4)'
    assert ws["T4"].value == "=((K4*$Q$2)*$D$16)"
    assert ws["U4"].value == "=Q4+S4+T4"
    assert ws["Y13"].value == "=$Y$11*X13" and ws["Z16"].value == "=X16*$Z$11"
    assert ws["AA22"].value == "=(Z22+Y15)-AA11"                      # the control cell
    assert [ws.cell(row=r, column=23).value for r in (13, 14, 16, 17, 21)] == [
        "KAR 1 Pablo", "KAR 2 Edwin", "Buyer and Cost Manager (Lupita)", "Logistics Supervisor (Sara Reyes)",
        "Adm. 2 Monica"]


def test_only_paid_lines_are_distributed():
    ws, _ = sheet()
    assert ws["Y5"].value.startswith('=IF(P5="Paid",U5*') and ws["Z5"].value.startswith('=IF(P5="Paid",U5*')
    assert [ws[f"P{r}"].value for r in range(4, 11)] == ["Paid", "Pending", None, "Paid", None, "Paid", None]


def test_the_engine_gives_the_templates_numbers_for_the_same_inputs():
    lines = [E.Line(job=x.file, billed=x.billed or 0, cost=x.cost or 0, paid={"Paid": True, "Pending": False}.get(x.payment))
             for x in T.SAMPLE_ROWS]
    bookings = sum(x.sale_price or 0 for x in T.SAMPLE_ROWS)
    m = E.calc_month("s", lines, bookings, E.Settings(budgets={"s": T.SAMPLE_BUDGET}, margin_waits_for_cost=False))
    assert bookings == 840000 and m.reach == pytest.approx(1.008)
    assert m.total == pytest.approx(8822.16) and m.payable_total == pytest.approx(5091.90)
    assert [round(x.total, 2) for x in m.lines] == [3590.40, 3730.26, 0, 1817.70, 0, -316.20, 0]


def test_what_is_not_known_is_left_empty():
    ws, info = sheet([T.Row("110001", billed=1000.0, cost=400.0)], budget=5000.0, bookings=5000.0)
    assert [ws.cell(row=4, column=c).value for c in (1, 3, 4, 5, 13, 16)] == [None] * 6
    assert ws["D5"].value == 5000.0 and ws["D5"].comment is not None   # bookings typed in, not summed
    assert ws["D5"].fill.fgColor.rgb.endswith("FFFF00")
    assert ws["I5"].value == "=IF(H5=0,0,D5/H5)"


def test_a_billed_file_with_no_cost_yet_says_pending_and_has_no_margin_part():
    ws, _ = sheet([T.Row("110001", billed=1000.0, cost_pending=True), T.Row("110002", billed=-500.0, cost=0.0)],
                  budget=5000.0, bookings=5000.0)
    assert ws["L4"].value == T.PENDING and ws["L5"].value == 0.0
    assert ws["N4"].value.startswith("=IF(OR(ISTEXT(L4)")             # text in the cost cell: no gross margin
    assert ws["L6"].value == "=SUM(L4:L5)"                             # a sum ignores the word


def test_rows_from_the_report_fill_only_the_real_block():
    rep = {"lines": [{"month": "2026-08", "job": "110010", "billed": 100.0, "cost": 60.0, "cost_posted": True},
                     {"month": "2026-08", "job": "110011", "billed": 50.0, "cost": 0.0, "cost_posted": False},
                     {"month": "2026-09", "job": "110012", "billed": 70.0, "cost": 10.0, "cost_posted": True}]}
    rows = T.rows_for_month(rep, "2026-08")
    assert [(x.month, x.file, x.billed, x.cost, x.cost_pending) for x in rows] == [
        ("August", "110010", 100.0, 60.0, False), ("", "110011", 50.0, None, True)]
    assert all(x.sale_price is None and x.status == "" and x.payment == "" for x in rows)
    assert T.rows_for_month(rep, "2026-10") == []


def test_target_margin_gets_its_own_cell_when_it_differs_from_the_weight():
    ws = openpyxl.Workbook().active
    T.write(ws, T.SAMPLE_ROWS, E.Settings(target_margin=0.35), T.SAMPLE_BUDGET)
    assert ws["F15"].value == 0.35 and ws["R4"].value == '=IF(O4="","",O4/$F$15)'
    assert ws["S4"].value == '=IF(OR(R4="",K4=0),0,((K4*$Q$2)*$D$15)*R4)'      # the weight stays D15


def _report():
    from commissions import report as R
    ln = lambda job, month, billed, cost, posted=True: {"month": month, "job": job, "billed": billed, "cost": cost,
                                                         "cost_posted": posted}
    mo = lambda month, rule, known, bookings=None: {"month": month, "rule": rule, "budget": S.budgets[month],
                                                     "bookings": bookings, "bookings_known": known}
    return {"months": [mo("2026-03", "Old calculation", False), mo("2026-04", "Old calculation", False),
                       mo("2026-05", "New calculation", True, 3_000_000.0)],
            "lines": [ln("110001", "2026-03", 1000.0, 400.0), ln("110002", "2026-04", 2000.0, 0.0, posted=False),
                      ln("110003", "2026-05", 3000.0, 900.0), ln("110004", "2026-05", -500.0, 0.0)]}


def test_whole_report_is_one_sheet_per_month_in_month_order():
    wb = T.build_book(_report(), S)
    assert wb.sheetnames == ["Mar 2026", "Apr 2026", "May 2026"] and wb.active.title == "May 2026"
    may = wb["May 2026"]
    assert [may[f"B{r}"].value for r in (4, 5)] == ["110003", "110004"] and may["A4"].value == "May"
    assert may["D6"].value == 3_000_000.0 and may["H6"].value == S.budgets["2026-05"]
    assert wb["Apr 2026"]["L4"].value == T.PENDING and wb["Apr 2026"]["D5"].value == S.budgets["2026-04"]


def test_each_sheet_says_what_is_missing_and_which_rule_the_month_follows():
    wb = T.build_book(_report(), S)
    notes = lambda ws: " ".join(str(c.value) for row in ws.iter_rows(min_col=2, max_col=2) for c in row if c.value)
    assert "paid on the old calculation" in notes(wb["Mar 2026"])
    assert "paid on the old calculation" in notes(wb["Apr 2026"]) and "no cost in the books yet" in notes(wb["Apr 2026"])
    may = notes(wb["May 2026"])
    assert "old calculation" not in may and "entered on the commissions page" in may and "16.50" in may
    assert "set equal to the budget" in notes(wb["Mar 2026"]) and "set equal to the budget" not in may


def test_a_report_with_no_month_still_opens():
    wb = T.build_book({"months": [], "lines": []}, S)
    assert wb.sheetnames == ["No months"]
