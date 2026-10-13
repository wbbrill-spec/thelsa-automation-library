"""The /commissions pages: who may open them, what an upload keeps, and that the
page, the stored report and the Excel download all show the same numbers."""
import io
import json

import openpyxl
import pytest
from flask import Flask

from commissions import template_sheet as T
from commissions import engine as E
from commissions import finance as F
from commissions import report as R
from commissions import store
from commissions import web as W
from test_commissions_finance import ROWS, book, cost, row

LUPITA, OUTSIDER = "maria.gonzalez@thelsa.com", "someone.else@thelsa.com"
NAMES = ("Acme", "Jane Roe", "Big Co", "Mary Major", "EMBAJADA", "Marcel", "MARCEL", "Supplier", "RECLASIF")


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("COMMISSION_DATABASE_URL", f"sqlite:///{tmp_path}/c.db")
    monkeypatch.delenv("COMMISSION_USERS", raising=False)
    store.reset_for_tests()
    app = Flask(__name__)
    app.secret_key = "test"
    app.add_url_rule("/login", "login", lambda: "login")
    app.register_blueprint(W.commissions_bp)
    with app.test_client() as c:
        c.application_ref = app
        yield c
    store.reset_for_tests()


def login(c, email=LUPITA):
    with c.session_transaction() as s:
        s["user_email"] = email
        s["commission_csrf"] = "tok"


def xlsx(tmp_path, rows=ROWS, name="m.xlsx"):
    return open(book(tmp_path / name, rows), "rb").read()


def upload(c, data, filename="Margen x Expediente 2609.xlsx", csrf="tok"):
    return c.post("/commissions/upload", data={"csrf": csrf, "report": (io.BytesIO(data), filename)},
                  content_type="multipart/form-data")


# ── who may open it ──────────────────────────────────────────────────────────
def test_needs_a_login(client):
    assert client.get("/commissions").status_code == 302
    assert client.post("/commissions/upload").status_code == 302


def test_every_commission_route_is_closed_to_everyone_not_on_the_list(client):
    """A new route cannot ship open: every rule under /commissions is tried."""
    login(client, OUTSIDER)
    rules = [r for r in client.application_ref.url_map.iter_rules() if r.rule.startswith("/commissions")]
    assert len(rules) >= 4
    for r in rules:
        for method in r.methods - {"HEAD", "OPTIONS"}:
            assert client.open(r.rule, method=method).status_code == 404, (r.rule, method)


def test_the_list_is_lupita_mario_rogelio_and_bill_unless_render_says_otherwise(client, monkeypatch):
    assert W.allowed_users() == {"maria.gonzalez@thelsa.com", "mariotorres@thelsa.com", "rogeliofranco@thelsa.com",
                                 "bbrill@thelsa.com", "bill.brill@inflectionpointnow.com"}
    for who in ("MarioTorres@thelsa.com", "rogeliofranco@thelsa.com"):
        login(client, who)
        assert client.get("/commissions").status_code == 200
    monkeypatch.setenv("COMMISSION_USERS", "only.one@thelsa.com; ")
    login(client, LUPITA)
    assert client.get("/commissions").status_code == 404
    login(client, "only.one@thelsa.com")
    assert client.get("/commissions").status_code == 200


def test_page_before_any_upload(client):
    login(client)
    r = client.get("/commissions")
    assert r.status_code == 200 and b"No report has been uploaded yet" in r.data
    assert r.headers["Cache-Control"] == "no-store"
    assert client.get("/commissions/workbook").status_code == 302


# ── upload ───────────────────────────────────────────────────────────────────
def test_upload_needs_the_form_token(client, tmp_path):
    login(client)
    assert upload(client, xlsx(tmp_path), csrf="wrong").status_code == 400
    assert upload(client, xlsx(tmp_path), csrf="").status_code == 400
    assert store.get("finance_report")[0] is None


def test_upload_keeps_lines_and_no_names(client, tmp_path):
    login(client)
    assert upload(client, xlsx(tmp_path)).status_code == 302
    doc, meta = store.get("finance_report")
    assert meta["by"] == LUPITA and len(doc["records"]) == len(ROWS) - 1 and len(doc["sha256"]) == 64
    with store.engine().connect() as c:
        everything = " ".join(str(v) for row_ in c.exec_driver_sql("select * from commission_docs") for v in row_)
    page = client.get("/commissions").data.decode()
    for name in NAMES:
        assert name not in everything and name not in page
    assert "Report loaded: 19 lines dated 2026-01-12 to 2026-08-05" in page
    log = store.get("uploads")[0]
    assert len(log) == 1 and log[0]["by"] == LUPITA and log[0]["lines"] == 19


def test_page_shows_the_same_numbers_as_the_engine(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    page = client.get("/commissions").data.decode()
    rep = R.build(F.load_base(io.BytesIO(xlsx(tmp_path))), W.settings())
    assert f"{rep['total_new']['total']:,.2f}" in page and f"{rep['total_new']['billed']:,.0f}" in page
    assert "Trial. Not for payment." in page and "= budget" in page
    assert "Exchange rate looks mistyped" in page and "110633" in page
    assert "waits for cost" in page
    for p in rep["people"]:
        assert p["name"].replace("&", "&amp;") in page and f"{p['if_collected']:,.2f}" in page


@pytest.mark.parametrize("what", ["not excel name", "garbage", "money on a row with no job", "text amount", "empty"])
def test_a_report_that_cannot_be_read_changes_nothing(client, tmp_path, what):
    login(client)
    upload(client, xlsx(tmp_path))
    before = store.get("finance_report")[0]
    good = row("110001", 1000)
    stray = row(None, 500, typ=None, folio=2)
    stray[3] = None
    bad = {"not excel name": (xlsx(tmp_path), "report.csv"), "garbage": (b"PK\x03\x04 not a workbook", "r.xlsx"),
           "money on a row with no job": (xlsx(tmp_path, [good, stray], "s.xlsx"), "r.xlsx"),
           "text amount": (xlsx(tmp_path, [row("110001", "ver con Sr. Perez")], "t.xlsx"), "r.xlsx"),
           "empty": (b"", "r.xlsx")}[what]
    assert upload(client, *bad).status_code == 302
    page = client.get("/commissions").data.decode()
    assert store.get("finance_report")[0] == before
    assert len(store.get("uploads")[0]) == 1
    assert "Perez" not in page
    assert ("Nothing was changed" in page) or ("must be an Excel workbook" in page)
    if what in ("money on a row with no job", "text amount"):
        assert "row 4" in page or "row 3" in page


def test_a_second_upload_replaces_the_first_and_keeps_the_previous_one(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    first = store.get("finance_report")[0]
    upload(client, xlsx(tmp_path, [row("110001", 1000), cost("110001", 400)], "two.xlsx"))
    assert len(store.get("finance_report")[0]["records"]) == 2
    assert store.get("finance_report_previous")[0] == first
    assert len(store.get("uploads")[0]) == 2


def test_an_upload_too_large_is_refused(client, tmp_path, monkeypatch):
    login(client)
    monkeypatch.setattr(W, "MAX_UPLOAD", 100)
    upload(client, xlsx(tmp_path))
    assert store.get("finance_report")[0] is None
    assert "larger than 20 MB" in client.get("/commissions").data.decode()


def test_a_stored_report_that_was_changed_by_hand_is_refused(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    doc = store.get("finance_report")[0]
    doc["records"][0]["job"] = "ACME CORPORATION"
    store.put("finance_report", doc, by="x")
    page = client.get("/commissions").data.decode()
    assert "could not be read back" in page and "ACME" not in page
    assert client.get("/commissions/workbook").status_code == 302


# ── bookings ─────────────────────────────────────────────────────────────────
def test_bookings_change_sales_reach(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    budget = E.BUDGET_2026["2026-05"]
    r = client.post("/commissions/bookings", data={"csrf": "tok", "b_2026-05": f"{budget * 1.5:,.2f}", "b_2026-06": ""})
    assert r.status_code == 302
    assert store.get("bookings")[0] == {"2026-05": pytest.approx(round(budget * 1.5, 2))}
    page = client.get("/commissions").data.decode()
    assert "0.9000" in page and "150.0%" in page                 # May: 1.5 x 0.60
    assert "Bookings are not known for 2 month(s)" in page       # June and July of the fixture
    rep = R.build(F.from_records(store.get("finance_report")[0]["records"]), W.settings(), store.get("bookings")[0])
    assert rep["months"][0]["reach"] == pytest.approx(0.9) and f"{rep['total_new']['total']:,.2f}" in page


@pytest.mark.parametrize("bad", ["abc", "1.234.567,00", "-5", "0", "12,34", "1e6", "100.123"])
def test_bookings_that_are_not_an_amount_are_not_saved(client, bad):
    login(client)
    client.post("/commissions/bookings", data={"csrf": "tok", "b_2026-05": "2,000,000.00"})
    client.post("/commissions/bookings", data={"csrf": "tok", "b_2026-05": bad, "b_2026-06": "3,000,000"})
    assert store.get("bookings")[0] == {"2026-05": 2000000.0}
    assert "Bookings were not saved" in client.get("/commissions").data.decode()


def test_bookings_need_the_form_token_and_blank_clears_a_month(client):
    login(client)
    assert client.post("/commissions/bookings", data={"b_2026-05": "1"}).status_code == 400
    client.post("/commissions/bookings", data={"csrf": "tok", "b_2026-05": "2000000"})
    client.post("/commissions/bookings", data={"csrf": "tok", "b_2026-05": ""})
    assert store.get("bookings")[0] == {}


# ── Excel download ───────────────────────────────────────────────────────────
def test_workbook_download_matches_the_page(client, tmp_path):
    """The download is the report in Finance's own layout: one sheet per month."""
    login(client)
    upload(client, xlsx(tmp_path))
    client.post("/commissions/bookings", data={"csrf": "tok", "b_2026-05": "3,000,000.00"})
    r = client.get("/commissions/workbook")
    assert r.status_code == 200 and r.headers["Content-Disposition"].endswith('to 2026-08-05.xlsx"')
    assert r.headers["Cache-Control"] == "no-store"
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    rep = R.build(F.load_base(io.BytesIO(xlsx(tmp_path))), W.settings(), {"2026-05": 3000000.0})
    assert wb.sheetnames == [T.sheet_name(m["month"]) for m in rep["months"]] and "May 2026" in wb.sheetnames
    text = " ".join(str(c.value) for ws in wb for r_ in ws.iter_rows() for c in r_ if c.value is not None)
    for name in NAMES:
        assert name not in text
    jobs_in_book = []
    for m in rep["months"]:
        ws = wb[T.sheet_name(m["month"])]
        assert (ws["A2"].value, ws["K2"].value, ws["W2"].value) == ("BOOKED QUOTE", "REAL", "COMMISSION DISTRIBUTION")
        mine = [x for x in rep["lines"] if x["month"] == m["month"]]
        tot = 4 + len(mine)
        assert ws[f"A{tot}"].value == "Total Bookings" and ws[f"H{tot}"].value == m["budget"]
        # bookings typed on the page are used; a month without them keeps the budget as a placeholder
        assert ws[f"D{tot}"].value == (3000000.0 if m["month"] == "2026-05" else m["budget"])
        for i, x in enumerate(mine):
            row = 4 + i
            jobs_in_book.append(ws[f"B{row}"].value)
            assert ws[f"K{row}"].value == pytest.approx(x["billed"])
            assert ws[f"L{row}"].value == (pytest.approx(x["cost"]) if x["cost_posted"] else T.PENDING)
            assert ws[f"P{row}"].value is None and ws[f"U{row}"].value == f"=Q{row}+S{row}+T{row}"
    assert jobs_in_book == [x["job"] for x in rep["lines"]]


# ── the report as numbers ────────────────────────────────────────────────────
def test_report_totals_are_the_sum_of_their_months_and_people_get_all_of_it(tmp_path):
    lines = F.load_base(book(tmp_path / "r.xlsx", ROWS))
    rep = R.build(lines, W.settings(), {"2026-05": 1_000_000, "2026-07": float("nan"), "2026-06": -5, "x": "1"})
    assert rep["placeholder_months"] == ["2026-06", "2026-07"]       # nan and negative bookings are not bookings
    assert rep["total_all"]["total"] == pytest.approx(sum(m["total"] for m in rep["months"]))
    new_lines = [r for r in rep["lines"] if r["month"] >= R.NEW_FROM]
    assert 0 < len(new_lines) < len(rep["lines"])                    # the fixture has May (old) and later (new)
    assert rep["total_new"]["total"] == pytest.approx(sum(r["total"] for r in new_lines))
    assert sum(p["if_collected"] for p in rep["people"]) == pytest.approx(rep["total_new"]["total"])
    assert rep["control"] == pytest.approx(0, abs=1e-6)
    assert rep["total_new"]["payable"] == 0
    assert [c["kind"] for c in rep["checks"]][:2] == ["fx_outlier", "embassy_billed_in_scope"]
    assert json.dumps(rep)                                           # plain numbers and text all the way down


def test_a_month_without_a_budget_is_not_calculated_and_is_named(tmp_path):
    from test_commissions_finance import D
    lines = F.load_base(book(tmp_path / "r.xlsx", [row("110001", 1000), row("110002", 2000, date=D(2027, 1, 5), folio=2)]))
    rep = R.build(lines, W.settings())
    assert rep["months_without_budget"] == ["2027-01"] and [m["month"] for m in rep["months"]] == ["2026-05"]


def test_records_round_trip_and_refuse_anything_the_reader_did_not_write(tmp_path):
    lines = F.load_base(book(tmp_path / "r.xlsx", ROWS))
    recs = json.loads(json.dumps(F.to_records(lines)))
    assert F.from_records(recs) == lines
    for field, value in (("job", "ACME"), ("series", "GARCIA"), ("folio", "12b"), ("type", "PRIVADO"),
                         ("net_sales_mxn", "1000"), ("net_sales_mxn", float("nan")), ("cancels", "see Juan"),
                         ("currency", "GBP"), ("date", "yesterday"), ("extra", 1), ("entity", "Garcia SA")):
        bad = json.loads(json.dumps(recs))
        bad[3][field] = value
        with pytest.raises(F.FinanceReportError) as e:
            F.from_records(bad)
        assert "stored line 4" in str(e.value) and str(value) not in str(e.value)


# ── MXN / USD and EN / ES ────────────────────────────────────────────────────
def test_page_carries_both_languages_and_both_currencies(client, tmp_path):
    import re
    login(client)
    upload(client, xlsx(tmp_path))
    page = client.get("/commissions").data.decode()
    assert 'id="ccySeg"' in page and 'data-c="usd"' in page and 'id="langSeg"' in page and 'data-l="es"' in page
    for es in ("Reporte de comisiones TMS", "Prueba. No usar para pago.", "Facturado", "Costo real", "Margen bruto",
               "Alcance de ventas", "Por mes de factura", "Puntos por revisar",
               "Ventas contratadas por mes", "Subir reporte", "Descargar en Excel", "Cálculo nuevo",
               "El tipo de cambio parece mal capturado", "fila 17: tipo de cambio 12.6231", "espera costo",
               "Compradora y Gerente de Costos (Lupita)", "Reporte cargado: 19 líneas", "mayo 2026"):
        assert es in page, es
    rep = R.build(F.load_base(io.BytesIO(xlsx(tmp_path))), W.settings())
    spans = re.findall(r'class="m" data-x="([-0-9.]+)" data-u="([-0-9.]*)" data-d="(\d)">([^<]*)<', page)
    assert len(spans) > 100
    assert all(u != "" for _x, u, _d, _t in spans)                  # every amount has its dollar figure
    pairs = {(round(float(x), 2), round(float(u), 2)) for x, u, _d, _t in spans}
    new = rep["total_new"]
    assert (round(new["total"], 2), round(new["usd"]["total"], 2)) in pairs
    may = rep["months"][0]
    assert may["month"] == "2026-05" and may["rate"] == 16.5                      # the commission's own rate
    assert (round(may["billed"], 2), round(may["billed"] / 16.5, 2)) in pairs
    assert "internal rate of 16.50 pesos per dollar" in page and "tipo de cambio interno de 16.50" in page
    # without the script the page still reads in pesos
    assert f">{new['total']:,.2f}<" in page


def test_dollar_view_at_the_internal_rate_shows_the_dollars_invoiced():
    from test_commissions_finance import USD
    import tempfile, pathlib
    lines = F.load_base(book(pathlib.Path(tempfile.mkdtemp()) / "r.xlsx",
                             [row("110001", 175000, usd=10000, ccy=USD, fx=17.5), row("110002", 33000, folio=2)]))
    rep = R.build(lines, W.settings())
    assert W.settings().sales_fx == "internal"
    assert rep["rates"] == {"by_month": {"2026-05": 16.5}, "other": 16.5, "from_report": False}
    by = {x["job"]: x for x in rep["lines"]}
    assert by["110001"]["billed"] == pytest.approx(165000) and by["110001"]["usd"]["billed"] == pytest.approx(10000)
    assert by["110002"]["usd"]["billed"] == pytest.approx(2000)
    assert rep["total_all"]["usd"]["total"] == pytest.approx(rep["total_all"]["total"] / 16.5)


def test_dollar_view_with_spot_uses_each_months_own_rate_and_adds_up():
    from test_commissions_finance import D, USD
    rows = [row("110001", 175000, usd=10000, ccy=USD, fx=17.5), row("110002", 12623.1, usd=1000, ccy=USD, fx=12.6231, folio=2),
            row("110003", 176000, usd=10000, ccy=USD, fx=17.6, folio=3),
            row("110004", 180000, usd=10000, ccy=USD, fx=18.0, folio=4, date=D(2026, 6, 3)),
            row("110005", 90000, folio=5, date=D(2026, 7, 3))]                    # a month with no dollar invoice
    import tempfile, pathlib
    lines = F.load_base(book(pathlib.Path(tempfile.mkdtemp()) / "r.xlsx", rows))
    spot = E.Settings(budgets=dict(E.BUDGET_2026), sales_fx="spot")
    rep = R.build(lines, spot)
    rates = rep["rates"]["by_month"]
    assert rates["2026-05"] == pytest.approx(17.5) and rates["2026-06"] == pytest.approx(18.0)
    assert rates["2026-07"] == pytest.approx(17.55) == rep["rates"]["other"]      # the middle of all of them
    for m in rep["months"]:
        assert m["usd"]["billed"] == pytest.approx(m["billed"] / m["rate"])
    for k in R.MONEY:
        assert rep["total_all"]["usd"][k] == pytest.approx(sum(m["usd"][k] for m in rep["months"]))
    assert sum(p["usd"]["if_collected"] for p in rep["people"]) == pytest.approx(rep["total_new"]["usd"]["total"])
    assert sum(x["usd"]["total"] for x in rep["lines"]) == pytest.approx(rep["total_all"]["usd"]["total"])
    none = R.build([ln for ln in lines if ln.job == "110005"], spot)
    assert none["rates"]["from_report"] is False and none["rates"]["other"] == 16.5


def test_every_point_and_every_label_has_its_spanish(tmp_path):
    assert set(R.LABELS_ES) == set(R.LABELS)
    assert all(R.LABELS_ES[k] != R.LABELS[k] for k in R.LABELS)
    lines = F.load_base(book(tmp_path / "r.xlsx", ROWS))
    for s in (E.Settings(), E.Settings(sales_fx="internal")):
        for p in F.problems(lines, s):
            assert p["detail_es"] and p["detail_es"] != p["detail"], p["kind"]
            assert "row " not in p["detail_es"]


def test_messages_come_in_both_languages(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path, [row("110001", "ver nota")], "t.xlsx"))
    page = client.get("/commissions").data.decode()
    assert "the cell holds text, not an amount" in page and "la celda tiene texto, no un importe" in page
    assert "No se cambió nada" in page
    client.post("/commissions/bookings", data={"csrf": "tok", "b_2026-05": "abc"})
    page = client.get("/commissions").data.decode()
    assert "May 2026" in page and "No se guardaron las ventas contratadas" in page and "mayo 2026" in page



# ── the tile on the library home page ─────────────────────────────────────────
def _home():
    import pathlib
    return (pathlib.Path(__file__).parent / "index.html").read_text()


def test_home_tile_is_shown_only_to_the_people_who_may_open_the_report(monkeypatch):
    monkeypatch.delenv("COMMISSION_USERS", raising=False)
    home = _home()
    assert 'href="/commissions"' in home
    for email in W.DEFAULT_USERS:
        assert 'href="/commissions"' in W.home_tile(home, email)
        assert 'href="/commissions"' in W.home_tile(home, "  " + email.upper() + " ")
    for email in ("someone@thelsa.com", "", None):
        page = W.home_tile(home, email)
        assert "/commissions" not in page and "Commission" not in page
        assert "COMMISSION_TILE" not in page
        # the rest of the page is untouched
        assert "TMS Executive Dashboard" in page and "TIM Executive Dashboard" in page
        assert page.count("<div") == page.count("</div>")


def test_home_tile_follows_the_render_setting(monkeypatch):
    monkeypatch.setenv("COMMISSION_USERS", "new.person@thelsa.com")
    home = _home()
    assert 'href="/commissions"' in W.home_tile(home, "new.person@thelsa.com")
    assert "/commissions" not in W.home_tile(home, "bbrill@thelsa.com")


def test_the_new_calculation_starts_in_june_2026():
    """Lupita, 12 Oct 2026: April and May are old calculation."""
    assert R.NEW_FROM == "2026-06"
    assert [R.rule_for(m) for m in ("2026-01", "2026-04", "2026-05", "2026-06", "2026-12")] == [
        "Old calculation", "Old calculation", "Old calculation", "New calculation", "New calculation"]
    assert set(R.rule_for(f"2026-{m:02d}") for m in range(1, 13)) <= set(R.RULES_ES)


# ── the two Moveware reports ─────────────────────────────────────────────────
def _moveware_files(tmp_path):
    import test_commissions_quotes as TQ
    jobs = [TQ.J(j["job"], ex=250000.0) for j in ({"job": x} for x in sorted({r[2] for r in ROWS[:-1] if r[1] == "CORP-PART"}))]
    rem = open(TQ.xl(tmp_path, TQ.removals_rows(jobs), "rem.xlsx"), "rb").read()
    cst = open(TQ.xl(tmp_path, TQ.costing_rows([(jobs[0]["job"], "IMP", "Private", [TQ.O()])]), "cst.xlsx"), "rb").read()
    return rem, cst, TQ.NAMES


def _send(c, route, data, csrf="tok", filename="report.xlsx"):
    return c.post(route, data={"csrf": csrf, "report": (io.BytesIO(data), filename)}, content_type="multipart/form-data")


def test_moveware_reports_are_loaded_without_their_names_and_fill_the_quotes(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    before = client.get("/commissions").data.decode()
    assert "Removals Summary: not loaded yet." in before and "Move Costing Sheet: not loaded yet." in before
    rem, cst, names = _moveware_files(tmp_path)
    assert _send(client, "/commissions/upload-removals", rem).status_code == 302
    page = client.get("/commissions").data.decode()
    assert "Removals Summary loaded:" in page and "Customer names in it were not kept." in page
    assert _send(client, "/commissions/upload-costing", cst).status_code == 302
    page = client.get("/commissions").data.decode()
    assert "Move Costing Sheet loaded: 1 jobs." in page and "Removals Summary in use: uploaded" in page
    assert "2026-01-01" in page and "2026-10-09" in page and "Files first invoiced" in page and "the booked quotes" in page
    with store.engine().connect() as c:
        everything = " ".join(str(v) for row_ in c.exec_driver_sql("select * from commission_docs") for v in row_)
    for name in names + tuple(NAMES):
        assert name not in everything and name not in page
    # the page, the stored report and the Excel download agree
    rep = W._report(F.from_records(store.get("finance_report")[0]["records"]), W.settings())
    assert rep["quotes_loaded"] and any(m["bookings_source"] == "quotes" for m in rep["months"])
    assert f"{rep['total_new']['total']:,.2f}" in page
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/commissions/workbook").data))
    first = [x for x in rep["lines"] if x["sale_price"] is not None][0]
    ws = wb[T.sheet_name(first["month"])]
    row_ = next(r_ for r_ in range(4, 40) if ws[f"B{r_}"].value == first["job"])
    assert (ws[f"C{row_}"].value, ws[f"D{row_}"].value) == ("Ganado", pytest.approx(first["sale_price"]))
    log = store.get("uploads")[0]
    assert [x["what"] for x in log] == ["finance_report", "removals_summary", "costing_sheet"]


def test_moveware_uploads_need_the_token_and_refuse_other_files(client, tmp_path):
    login(client)
    rem, cst, _ = _moveware_files(tmp_path)
    for route in ("/commissions/upload-removals", "/commissions/upload-costing"):
        assert _send(client, route, rem, csrf="wrong").status_code == 400
        assert _send(client, route, b"PK\x03\x04 not a workbook").status_code == 302
        assert "Nothing was changed" in client.get("/commissions").data.decode()
        assert _send(client, route, rem, filename="report.csv").status_code == 302
        assert "must be an Excel workbook" in client.get("/commissions").data.decode()
    assert _send(client, "/commissions/upload-removals", cst).status_code == 302      # the wrong report in each box
    assert "does not look like the Moveware Removals Summary" in client.get("/commissions").data.decode()
    assert _send(client, "/commissions/upload-costing", rem).status_code == 302
    assert "does not look like the Moveware Move Costing Sheet" in client.get("/commissions").data.decode()
    assert store.get("removals_summary")[0] is None and store.get("costing_sheet")[0] is None


def test_the_page_works_with_a_stored_moveware_report_that_cannot_be_read_back(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    store.put("removals_summary", {"records": {"jobs": [{"job": "x"}], "control": {}}}, by=LUPITA)
    r = client.get("/commissions")
    assert r.status_code == 200 and b"could not be read back" in r.data and b"Removals Summary: not loaded yet." in r.data
    assert client.get("/commissions/workbook").status_code == 200


# ── payment status ───────────────────────────────────────────────────────────
def _lines_of(client):
    return W._report(F.from_records(store.get("finance_report")[0]["records"]), W.settings())


def test_only_lines_marked_paid_are_payable(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    rep = _lines_of(client)
    assert rep["total_all"]["payable"] == 0 and rep["paid_lines"] == 0
    assert "No line is marked as paid yet" in client.get("/commissions").data.decode()
    a, b, c = rep["lines"][0], rep["lines"][1], rep["lines"][-1]
    r = client.post("/commissions/payments", data={"csrf": "tok", f"p_{a['month']}_{a['job']}": "Paid",
                                                   f"p_{b['month']}_{b['job']}": "Pending", f"p_{c['month']}_{c['job']}": "Paid"})
    assert r.status_code == 302 and store.get("payments")[1]["by"] == LUPITA
    rep = _lines_of(client)
    by = {(x["month"], x["job"]): x for x in rep["lines"]}
    assert [by[(x["month"], x["job"])]["payment"] for x in (a, b, c)] == ["Paid", "Pending", "Paid"]
    assert rep["paid_lines"] == 2 and rep["pending_lines"] == 1
    assert rep["total_all"]["payable"] == pytest.approx(a["total"] + c["total"])
    assert rep["total_all"]["payable"] + rep["total_all"]["held"] == pytest.approx(rep["total_all"]["total"])
    assert sum(p["payable"] for p in rep["people"]) == pytest.approx(rep["total_new"]["payable"])
    page = client.get("/commissions").data.decode()
    assert "Payment status saved: 3 line(s) changed, 2 marked Paid in all." in page
    assert f'name="p_{a["month"]}_{a["job"]}"' in page and "2 " in page and "are paid. Only those are payable." in page
    # the Excel sheet carries the status, so Finance's own formulas distribute the same lines
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/commissions/workbook").data))
    ws = wb[T.sheet_name(a["month"])]
    row_ = next(r_ for r_ in range(4, 40) if ws[f"B{r_}"].value == a["job"])
    assert ws[f"P{row_}"].value == "Paid" and ws[f"Y{row_}"].value.startswith(f'=IF(P{row_}="Paid"')
    # setting a line back to "not known" removes the mark; lines not in the form keep theirs
    client.post("/commissions/payments", data={"csrf": "tok", f"p_{a['month']}_{a['job']}": ""})
    assert set(store.get("payments")[0]) == {f"{b['month']}|{b['job']}", f"{c['month']}|{c['job']}"}


def test_a_pasted_list_of_files_marks_one_month(client, tmp_path):
    login(client)
    upload(client, xlsx(tmp_path))
    rep = _lines_of(client)
    month = rep["lines"][0]["month"]
    jobs = [x["job"] for x in rep["lines"] if x["month"] == month]
    text = f"{jobs[0]}\n file {jobs[1].lower()}, 999999 and Jane Roe 12"
    r = client.post("/commissions/payments", data={"csrf": "tok", "bulk_jobs": text, "bulk_month": month, "bulk_status": "Paid"})
    assert r.status_code == 302
    assert store.get("payments")[0] == {f"{month}|{jobs[0]}": "Paid", f"{month}|{jobs[1]}": "Paid"}
    page = client.get("/commissions").data.decode()
    assert "2 line(s) of May 2026 marked Paid. Not in May 2026, left out: 999999." in page and "Jane Roe" not in page
    for bad in ({"bulk_month": "2026-13"}, {"bulk_status": "Maybe"}, {"bulk_month": ""}):
        data = {"csrf": "tok", "bulk_jobs": jobs[2] if len(jobs) > 2 else jobs[0], "bulk_month": month, "bulk_status": "Pending", **bad}
        client.post("/commissions/payments", data=data)
        assert "Nothing was changed" in client.get("/commissions").data.decode()
    assert len(store.get("payments")[0]) == 2


def test_payment_status_needs_the_token_a_report_and_plain_values(client, tmp_path):
    login(client)
    assert client.post("/commissions/payments", data={"p_2026-05_110001": "Paid"}).status_code == 400
    assert client.post("/commissions/payments", data={"csrf": "tok", "p_2026-05_110001": "Paid"}).status_code == 302
    assert store.get("payments")[0] is None                                   # no report loaded: nothing to mark
    upload(client, xlsx(tmp_path))
    client.post("/commissions/payments", data={"csrf": "tok", "p_2026-05_110001": "Yes", "p_2031-01_999999": "Paid",
                                               "p_2026-05_110002": "Paid"})
    assert store.get("payments")[0] == {"2026-05|110002": "Paid"}             # unknown words and unknown lines are ignored
    # marks left over from an older report are ignored by the calculation and named on the page
    store.put("payments", {"2026-05|110002": "Paid", "2025-01|123456": "Paid", "junk": "Paid", "2026-05|110001": "Sure"}, by=LUPITA)
    rep = _lines_of(client)
    assert rep["paid_lines"] == 1 and rep["payments_unused"] == ["2025-01|123456"]
    assert b"2025-01|123456" in client.get("/commissions").data
    assert R.clean_payments({"2026-05|110002": "Paid", "2026-05|110001": "Pending", "x": "Paid", "2026-5|1": "Paid"}) == {
        ("2026-05", "110002"): True, ("2026-05", "110001"): False}


# ── Finance's payments report ────────────────────────────────────────────────
def _payfile(tmp_path, rows, name="pay.xlsx", **kw):
    import test_commissions_payments as TP
    return open(TP.report(tmp_path, mxp=rows, name=name, **kw), "rb").read(), TP


def test_payments_report_makes_collected_lines_payable(client, tmp_path):
    import test_commissions_payments as TP
    login(client)
    upload(client, xlsx(tmp_path))
    assert b"No payments report loaded yet." in client.get("/commissions").data
    rows = [TP.pay(46150, 10, 203000.0, series="CORP"), TP.pay(46151, 77777, 5.0, series="CORP")]
    # the invoice is of TMS LLC: the same payment in another company's report is of another invoice
    other, _ = _payfile(tmp_path, rows, name="trs.xlsx")
    _send(client, "/commissions/upload-payments", other)
    assert "Payments report loaded for TRS" in client.get("/commissions").data.decode()
    assert {(x["month"], x["job"]): x["payment"] for x in _lines_of(client)["lines"]}[("2026-05", "110001")] == ""
    store.put("payment_reports", {"sources": {}}, by=LUPITA)
    data, _ = _payfile(tmp_path, rows, company="THELSA MOBILITY SOLUTIONS LLC")
    assert _send(client, "/commissions/upload-payments", data).status_code == 302
    page = client.get("/commissions").data.decode()
    assert "Payments report loaded for TMS LLC: 2 payments from 2026-01-01 to 2026-09-30." in page
    assert "Customer names and bank references in it were not kept." in page and "CORP 77777" in page
    rep = _lines_of(client)
    by = {(x["month"], x["job"]): x for x in rep["lines"]}
    assert (by[("2026-05", "110001")]["payment"], by[("2026-05", "110001")]["payment_source"]) == ("Paid", "file")
    assert rep["total_all"]["payable"] == pytest.approx(sum(x["total"] for x in rep["lines"] if x["payment"] == "Paid")) != 0
    with store.engine().connect() as c:
        everything = " ".join(str(v) for row_ in c.exec_driver_sql("select * from commission_docs") for v in row_)
    for name in TP.NAMES + tuple(NAMES):
        assert name not in everything and name not in page
    # the Excel sheet marks the same line, and a mark by hand on the page still wins
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/commissions/workbook").data))
    ws = wb["May 2026"]
    row_ = next(r_ for r_ in range(4, 40) if ws[f"B{r_}"].value == "110001")
    assert ws[f"P{row_}"].value == "Paid"
    client.post("/commissions/payments", data={"csrf": "tok", "p_2026-05_110001": "Pending"})
    assert {(x["month"], x["job"]): x["payment"] for x in _lines_of(client)["lines"]}[("2026-05", "110001")] == "Pending"
    # a new report for the same company replaces the old one; another company's report is kept beside it
    data2, _ = _payfile(tmp_path, [TP.pay(46150, 10, 1.0, series="CORP")], name="b.xlsx", company="Thelsa Mobility Solutions L.L.C.")
    _send(client, "/commissions/upload-payments", data2)
    data3, _ = _payfile(tmp_path, [TP.pay(46150, 11, 116000.0, series="PART")], name="c.xlsx", company="Thelsa Mobility Solutions, S.A. de C.V.")
    _send(client, "/commissions/upload-payments", data3)
    data4, _ = _payfile(tmp_path, [TP.pay(46150, 12, 5.0, series="PART")], name="d.xlsx", company="Some Other Company")
    _send(client, "/commissions/upload-payments", data4)
    assert "a company this page does not know by name" in client.get("/commissions").data.decode()
    doc = store.get("payment_reports")[0]
    assert len(doc["sources"]) == 3 and {"TMS", "TMS LLC"} < set(doc["sources"]) and "Some Other" not in json.dumps(doc)
    assert len(doc["sources"]["TMS LLC"]["records"]["payments"]) == 1
    assert len(W._collections()["sources"]) == 3
    assert W._collections()["error"] is None


def test_payments_upload_needs_the_token_and_refuses_other_files(client, tmp_path):
    import test_commissions_payments as TP
    login(client)
    data, _ = _payfile(tmp_path, [TP.pay(46150, 10, 203000.0, series="CORP")])
    assert _send(client, "/commissions/upload-payments", data, csrf="nope").status_code == 400
    assert _send(client, "/commissions/upload-payments", xlsx(tmp_path)).status_code == 302          # Finance's margin report
    assert "does not look like CONTPAQ" in client.get("/commissions").data.decode()
    assert _send(client, "/commissions/upload-payments", b"junk").status_code == 302
    assert "Nothing was changed" in client.get("/commissions").data.decode()
    assert store.get("payment_reports")[0] is None
    store.put("payment_reports", {"sources": {"TRS": {"records": {"payments": [{"x": 1}], "info": {}}}}}, by=LUPITA)
    upload(client, xlsx(tmp_path))
    r = client.get("/commissions")
    assert r.status_code == 200 and b"stored payments report could not be read back" in r.data


# ── Finance's invoice register ───────────────────────────────────────────────
def test_invoice_register_is_loaded_without_names_and_shows_whose_reports_are_needed(client, tmp_path):
    import test_commissions_invoices as TI
    import test_commissions_payments as TP
    login(client)
    upload(client, xlsx(tmp_path))
    assert "Invoice register: not loaded yet." in client.get("/commissions").data.decode()
    data = open(TI.register(tmp_path, [TI.inv("LLC", "CORP", 110001, TI.D(2026, 5, 4), "CORP", 10, 10000.0, 10000.0, ccy="USD"),
                                       TI.inv("SA", "PART", 110002, TI.D(2026, 5, 9), "PART", 11, 116000.0, 100000.0)]), "rb").read()
    assert _send(client, "/commissions/upload-invoices", data).status_code == 302
    page = client.get("/commissions").data.decode()
    assert "Invoice register loaded: 2 invoices from 2026-05-04 to 2026-05-09." in page and "Invoice register in use:" in page
    assert "Company that issued the invoice" in page and "not loaded" in page          # no payments report yet
    with store.engine().connect() as c:
        everything = " ".join(str(v) for row_ in c.exec_driver_sql("select * from commission_docs") for v in row_)
    for name in TI.NAMES + tuple(NAMES):
        assert name not in everything and name not in page
    # 110002's invoice carries IVA: paid without it, the line is not collected; in full, it is
    for amount, want in ((100000.0, "Pending"), (116000.0, "Paid")):
        pay_, _ = _payfile(tmp_path, [TP.pay(46160, 11, amount, series="PART")], company="Thelsa Mobility Solutions SA de CV")
        _send(client, "/commissions/upload-payments", pay_)
        assert {x["job"]: x["payment"] for x in _lines_of(client)["lines"] if x["month"] == "2026-05"}["110002"] == want
    assert client.get("/commissions").status_code == 200


def test_invoice_register_upload_needs_the_token_and_refuses_other_files(client, tmp_path):
    login(client)
    assert _send(client, "/commissions/upload-invoices", xlsx(tmp_path), csrf="nope").status_code == 400
    assert _send(client, "/commissions/upload-invoices", xlsx(tmp_path)).status_code == 302          # Finance's margin report
    assert "does not look like Finance" in client.get("/commissions").data.decode()
    assert _send(client, "/commissions/upload-invoices", b"junk").status_code == 302
    assert "Nothing was changed" in client.get("/commissions").data.decode() and store.get("invoice_register")[0] is None
    store.put("invoice_register", {"records": {"invoices": [{"x": 1}]}}, by=LUPITA)
    upload(client, xlsx(tmp_path))
    r = client.get("/commissions")
    assert r.status_code == 200 and "The stored invoice register could not be read back" in r.data.decode()
    client.get("/logout", follow_redirects=False)
    login(client, OUTSIDER)
    assert _send(client, "/commissions/upload-invoices", b"x").status_code == 404
